import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel

from config import field_mappings_public, optional_field_keys
from excel_layout import PUBLICATION_SHEET_NAME
from settings import (
    DATE_CELL_FORMAT,
    EXCEL_PATH,
    FULL_CONTRACT_PRICE_FIELD,
    FULL_EIS_CONTRACT_PRICE_FIELD,
    HEADER_ROW,
    NOTICE_COLUMN,
)


NUMERIC_PUBLIC_FIELDS = {
    "Цена контракта",
    "Цена контракта ЕИС",
    "Цена по протоколу",
    "Доля ЛП отечественного производства, %",
    "Доля ЛП иностранного производства, %",
    "Общее количество товара",
    "Количество товара по I этапу",
    "Количество товара по II этапу",
    "Количество товара по III этапу",
    "Количество товара по IV этапу",
    "Количество товара по V этапу",
    "Количество единиц измерения Товара во вторичной (потребительской) упаковке",
}


@dataclass
class CellChange:
    sheet_name: str
    cell: str
    field: str
    old_value: object
    new_value: object
    kind: str = "CHANGE"


@dataclass
class WriteResult:
    changed: bool
    changes: list[CellChange]
    warnings: list[str]


def extract_notice_number(value):
    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()

    match = re.search(r'\d{19}', text)
    if match:
        return match.group(0)

    digits = re.sub(r'\D', '', text)
    if len(digits) == 18:
        return digits.zfill(19)

    return None


def normalize_header(value):
    return ' '.join(str(value or '').replace('\n', ' ').split()).lower()


def build_header_column_map(sheet, header_row=HEADER_ROW):
    header_columns = {}

    for cell in sheet[header_row]:
        if cell.value not in (None, ''):
            header_columns[normalize_header(cell.value)] = get_column_letter(cell.column)

    return header_columns


def resolve_column(sheet, column_or_header, header_columns=None):
    column_or_header = str(column_or_header).strip()

    if re.fullmatch(r'[A-Z]{1,3}', column_or_header):
        return column_or_header

    if header_columns is None:
        header_columns = build_header_column_map(sheet)

    normalized_name = normalize_header(column_or_header)
    column = header_columns.get(normalized_name)

    if not column:
        raise ValueError(f'Колонка не найдена в Excel: {column_or_header}')

    return column


def resolve_public_column(sheet, field_key, column_or_header, header_columns=None):
    dynamic_column = resolve_dynamic_public_column(sheet, field_key)
    if dynamic_column:
        return dynamic_column

    return resolve_column(sheet, column_or_header, header_columns)


def resolve_dynamic_public_column(sheet, field_key):
    stage = extract_stage_from_field(field_key)
    if not stage:
        return None

    if field_key.startswith("Количество товара по "):
        return find_stage_header_column(sheet, stage, ("эт",))

    if field_key.startswith("Срок поставки по "):
        return find_stage_header_column(sheet, stage, ("срок", "поставк"))

    if field_key.startswith("Срок предоставления документов по "):
        return find_stage_header_column(sheet, stage, ("срок", "предостав", "документ"))

    return None


def extract_stage_from_field(field_key):
    match = re.search(r"\b(I{1,3}|IV|V)\s+этап", field_key, flags=re.IGNORECASE)
    if not match:
        return None

    return match.group(1).upper()


def find_stage_header_column(sheet, stage, required_parts):
    stage_pattern = re.compile(rf"(^|\s){re.escape(stage.lower())}\s+эт(?:ап)?(\s|$)")

    for cell in sheet[HEADER_ROW]:
        if cell.value in (None, ""):
            continue

        header = normalize_header(cell.value)
        if all(part in header for part in required_parts) and stage_pattern.search(header):
            return get_column_letter(cell.column)

    return None


def iter_notice_rows(sheet, notice_column=NOTICE_COLUMN, start_row=2):
    notice_column = resolve_column(sheet, notice_column)

    for row in range(start_row, sheet.max_row + 1):
        notice_number = extract_notice_number(sheet[f'{notice_column}{row}'].value)
        if notice_number:
            yield row, notice_number


def is_date_field(field_key):
    return field_key.startswith("Срок") or "Дата" in field_key


def normalize_excel_compare_value(value, field_key):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None

        if is_date_field(field_key):
            date_match = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', value)
            if date_match:
                day, month, year = map(int, date_match.groups())
                return date(year, month, day)

            date_match = re.search(r'(\d{4})-(\d{2})-(\d{2})', value)
            if date_match:
                year, month, day = map(int, date_match.groups())
                return date(year, month, day)

        return value

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if is_date_field(field_key) and isinstance(value, (int, float)):
        try:
            return from_excel(value).date()
        except Exception:
            return value

    return value


def values_are_equal(old_value, new_value, field_key):
    old_value = normalize_excel_compare_value(old_value, field_key)
    new_value = normalize_excel_compare_value(new_value, field_key)

    if old_value is None and new_value is None:
        return True
    if old_value is None or new_value is None:
        return False

    if isinstance(old_value, date) or isinstance(new_value, date):
        return old_value == new_value

    old_identifier = get_long_digit_identifier(old_value)
    new_identifier = get_long_digit_identifier(new_value)
    if old_identifier is not None or new_identifier is not None:
        return old_identifier == new_identifier

    try:
        return abs(to_compare_number(old_value) - to_compare_number(new_value)) < 0.01
    except (TypeError, ValueError):
        return str(old_value).strip() == str(new_value).strip()


def get_long_digit_identifier(value):
    if not isinstance(value, str):
        return None

    text = value.strip()
    if not re.fullmatch(r'[\d\s\r\n]+', text):
        return None

    digits = re.sub(r'\D', '', text)
    if len(digits) < 12:
        return None

    return digits


def to_compare_number(value):
    if isinstance(value, str):
        value = value.replace(" ", "").replace("\u202f", "").replace(",", ".")

    return float(value)


def get_number_value(data, field_key):
    value = data.get(field_key)
    if value in (None, ""):
        return None

    try:
        return to_compare_number(value)
    except (TypeError, ValueError):
        return None


def normalize_public_write_value(field_key, value):
    if field_key not in NUMERIC_PUBLIC_FIELDS or not isinstance(value, str):
        return value

    text = value.strip()
    if not text:
        return value

    normalized = (
        text
        .replace(" ", "")
        .replace("\u00a0", "")
        .replace("\u202f", "")
        .replace(",", ".")
    )
    if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", normalized):
        return value

    number = float(normalized)
    if number.is_integer():
        return int(number)

    return number


def is_multiyear_publication_data(data):
    years = set(re.findall(r"\b20\d{2}\b", str(data.get("Период поставок по годам") or "")))
    return len(years) > 1


def find_package_quantity_column(sheet):
    for cell in sheet[HEADER_ROW]:
        header = normalize_header(cell.value)
        if (
            "количество единиц измерения" in header
            and "вторичной упаковке" in header
        ):
            return get_column_letter(cell.column)

    return None


def parse_package_multiplier(value):
    normalized_value = normalize_public_write_value(
        "Количество единиц измерения Товара во вторичной (потребительской) упаковке",
        value,
    )
    if isinstance(normalized_value, (int, float)) and not isinstance(normalized_value, bool):
        return normalized_value

    text = str(value or "").strip()
    if not text:
        return None

    numbers = []
    for part in re.split(r"[;\n]+", text):
        part = re.sub(r"^\s*\d+\.\s+", "", part.strip())
        part = part.rstrip(".").strip()
        number = normalize_public_write_value(
            "Количество единиц измерения Товара во вторичной (потребительской) упаковке",
            part,
        )
        if not isinstance(number, (int, float)) or isinstance(number, bool):
            return None
        numbers.append(number)

    if numbers and all(abs(float(number) - float(numbers[0])) < 1e-9 for number in numbers):
        return numbers[0]

    return None


def format_formula_number(value):
    if isinstance(value, int):
        return str(value)

    return format(float(value), ".15g")


def build_price_calculation_values(sheet, data, row):
    stage_end = "AT"
    stage_v = data.get("Количество товара по V этапу")
    if stage_v not in (None, ""):
        stage_v_column = find_stage_header_column(sheet, "V", ("эт",))
        if stage_v_column:
            stage_end = stage_v_column

    values = {
        "AP": f"=SUM(AQ{row}:{stage_end}{row})",
        "AM": f"=AG{row}/AP{row}",
        "AN": f"=AF{row}/AP{row}",
    }

    eis_price = get_number_value(data, "Цена контракта ЕИС")
    if eis_price is None:
        values["AF"] = f"=AE{row}"

    if is_multiyear_publication_data(data):
        full_price = get_number_value(data, FULL_EIS_CONTRACT_PRICE_FIELD)
        if full_price is None:
            full_price = get_number_value(data, FULL_CONTRACT_PRICE_FIELD)
        if full_price is not None:
            values["AG"] = normalize_public_write_value(
                "Цена контракта",
                str(full_price),
            )
    else:
        values["AG"] = f"=AF{row}"

    package_column = find_package_quantity_column(sheet)
    package_value = data.get(
        "Количество единиц измерения Товара во вторичной (потребительской) упаковке"
    )
    package_multiplier = parse_package_multiplier(package_value)
    if package_column and package_multiplier is not None:
        simple_package_value = normalize_public_write_value(
            "Количество единиц измерения Товара во вторичной (потребительской) упаковке",
            package_value,
        )
        if isinstance(simple_package_value, (int, float)) and not isinstance(simple_package_value, bool):
            values["AO"] = f"=AN{row}*{package_column}{row}"
        else:
            values["AO"] = f"=AN{row}*{format_formula_number(package_multiplier)}"

    return values


def write_price_calculation_block(sheet, data, row, apply_changes=True, sheet_name=None):
    changes = []
    warnings = []
    sheet_name = sheet_name or sheet.title
    calculation_values = build_price_calculation_values(sheet, data, row)

    for column, new_value in calculation_values.items():
        cell = sheet[f"{column}{row}"]
        if cell.value == new_value:
            continue

        changes.append(CellChange(
            sheet_name,
            cell.coordinate,
            f"Расчётная формула {column}",
            cell.value,
            new_value,
            kind="FORMULA",
        ))
        if apply_changes:
            cell.value = new_value

    if "AO" not in calculation_values:
        warnings.append(
            f"AO{row} Цена за упаковку: составное количество во вторичной упаковке "
            "не сводится к одному числу, исходное значение сохранено"
        )

    return WriteResult(changed=bool(changes), changes=changes, warnings=warnings)


def get_date_value(data, field_key):
    value = normalize_excel_compare_value(data.get(field_key), field_key)
    if isinstance(value, date):
        return value

    return None


def format_warn_value(value, field_key):
    value = normalize_excel_compare_value(value, field_key)

    if value is None:
        return "пусто"
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")

    return str(value)


def is_optional_date_field(field_key):
    return field_key in optional_field_keys and field_key.startswith("Срок")


def print_cell_change(notice_number, cell_address, field_key, old_value, new_value, prefix='CHANGE'):
    notice_prefix = f'{notice_number}: ' if notice_number else ''
    print(
        f'{prefix} {notice_prefix}{cell_address} {field_key}: '
        f'было {format_warn_value(old_value, field_key)}, '
        f'стало {format_warn_value(new_value, field_key)}'
    )


def should_print_cell_change(old_value, field_key):
    if field_key == "Цена контракта":
        try:
            if abs(to_compare_number(old_value)) < 0.01:
                return False
        except (TypeError, ValueError):
            pass

    return normalize_excel_compare_value(old_value, field_key) is not None


def get_compare_cell_value(cell, values_sheet=None):
    if isinstance(cell.value, str) and cell.value.startswith('='):
        if values_sheet is not None:
            cached_value = values_sheet[cell.coordinate].value
            if cached_value is not None:
                return cached_value

        return calculate_cell_formula(cell, values_sheet)

    return cell.value


def calculate_cell_formula(cell, values_sheet=None, seen=None):
    if seen is None:
        seen = set()

    if cell.coordinate in seen:
        return None

    seen.add(cell.coordinate)
    formula = cell.value

    sum_value = calculate_sum_formula(formula, cell.parent, values_sheet, seen)
    if sum_value is not None:
        return sum_value

    return calculate_simple_formula(formula)


def calculate_sum_formula(formula, sheet, values_sheet=None, seen=None):
    match = re.fullmatch(r'=\s*SUM\(([A-Z]{1,3}\d+):([A-Z]{1,3}\d+)\)\s*', formula, flags=re.IGNORECASE)
    if not match:
        return None

    start_cell, end_cell = match.groups()
    min_col, min_row, max_col, max_row = range_boundaries(f'{start_cell}:{end_cell}')
    total = 0
    found_value = False

    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            cell = sheet.cell(row, column)
            value = get_formula_reference_value(cell, values_sheet, seen)
            if value in (None, ''):
                continue

            try:
                total += to_compare_number(value)
                found_value = True
            except (TypeError, ValueError):
                continue

    if not found_value:
        return None

    if float(total).is_integer():
        return int(total)

    return total


def get_formula_reference_value(cell, values_sheet=None, seen=None):
    if values_sheet is not None:
        cached_value = values_sheet[cell.coordinate].value
        if cached_value is not None:
            return cached_value

    if isinstance(cell.value, str) and cell.value.startswith('='):
        return calculate_cell_formula(cell, values_sheet, seen)

    return cell.value


def calculate_simple_formula(formula):
    expression = formula.lstrip('=').replace(',', '.')
    if not re.fullmatch(r'[\d\s\u202f.+\-*/()]+', expression):
        return None

    try:
        return eval(expression, {"__builtins__": {}}, {})
    except Exception:
        return None


def write_data_to_row(
    sheet,
    data,
    row,
    notice_number=None,
    values_sheet=None,
    apply_changes=True,
    sheet_name=PUBLICATION_SHEET_NAME,
):
    header_columns = build_header_column_map(sheet)
    changed = False
    changes = []
    warnings = []

    for field_key, column_or_header in field_mappings_public:
        if field_key in data:
            try:
                column = resolve_public_column(sheet, field_key, column_or_header, header_columns)
            except ValueError as error:
                if field_key in optional_field_keys:
                    if data.get(field_key) in (None, ""):
                        continue
                    warnings.append(str(error))
                    continue
                raise
            cell = sheet[f'{column}{row}']
            if field_key == "Общее количество товара" and isinstance(cell.value, str) and cell.value.startswith('='):
                continue

            old_value = get_compare_cell_value(cell, values_sheet)
            new_value = normalize_public_write_value(field_key, data[field_key])
            needs_numeric_type_fix = (
                field_key in NUMERIC_PUBLIC_FIELDS
                and isinstance(new_value, (int, float))
                and not isinstance(new_value, bool)
                and cell.data_type not in {"n", "f"}
            )

            has_formula = isinstance(cell.value, str) and cell.value.startswith('=')
            overwrite_eis_formula = (
                field_key == "Цена контракта ЕИС"
                and new_value not in (None, "")
            )
            if has_formula and not overwrite_eis_formula:
                continue

            if field_key in optional_field_keys:
                if values_are_equal(old_value, new_value, field_key) and not needs_numeric_type_fix:
                    if is_optional_date_field(field_key):
                        if cell.number_format != DATE_CELL_FORMAT:
                            changes.append(CellChange(
                                sheet_name,
                                cell.coordinate,
                                f'{field_key} (формат)',
                                cell.number_format,
                                DATE_CELL_FORMAT,
                                kind='FORMAT',
                            ))
                            if apply_changes:
                                cell.number_format = DATE_CELL_FORMAT
                            changed = True
                    continue

                changes.append(CellChange(sheet_name, cell.coordinate, field_key, old_value, new_value, kind='WARN'))
            elif values_are_equal(old_value, new_value, field_key) and not needs_numeric_type_fix:
                continue
            else:
                changes.append(CellChange(sheet_name, cell.coordinate, field_key, old_value, new_value))

            if apply_changes:
                sheet[f'{column}{row}'] = new_value
            changed = True
            if apply_changes and is_optional_date_field(field_key):
                sheet[f'{column}{row}'].number_format = DATE_CELL_FORMAT

    price_result = write_price_calculation_block(
        sheet,
        data,
        row,
        apply_changes=apply_changes,
        sheet_name=sheet_name,
    )
    changed = changed or price_result.changed
    changes.extend(price_result.changes)
    warnings.extend(price_result.warnings)

    return WriteResult(changed=changed, changes=changes, warnings=warnings)


def get_excel_lock_path(workbook_path=EXCEL_PATH):
    workbook_path = Path(workbook_path)
    return workbook_path.with_name(f'~${workbook_path.name}')


def ensure_excel_is_closed(workbook_path=EXCEL_PATH):
    lock_path = get_excel_lock_path(workbook_path)
    if lock_path.exists():
        raise RuntimeError(f'закрой Excel-файл перед запуском: найден {lock_path.name}')
