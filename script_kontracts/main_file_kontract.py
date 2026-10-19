import argparse
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook

from data_checks import (
    check_suspicious_data,
    count_excel_fields,
    get_missing_fields,
    validate_data,
)
from document_search import find_document_pair
from eis_client import fetch_eis_data, format_eis_error
from excel_tools import (
    CellChange,
    ensure_excel_is_closed,
    get_date_value,
    get_number_value,
    iter_notice_rows,
    values_are_equal,
    write_data_to_row,
)
from excel_layout import (
    JOURNAL_SHEET_NAME,
    find_execution_sheet,
    get_publication_sheets,
    validate_workbook_structure,
)
from execution_monthly import write_execution_by_month
from fas_documents import (
    FAS_APPROVAL_FIELD,
    FAS_NOTIFICATION_FIELD,
    get_fas_decision,
    validate_fas_decision,
)
from journal_register import write_journal_register
from logging_utils import TeeOutput, create_log_path
from parser_docx import parse_document_pair
from public_sync import (
    apply_publication_changed_filter,
    copy_main_publication_to_shapka,
)
from settings import (
    BASE_DIR,
    EXCEL_PATH,
    FAS_PRICE_THRESHOLD,
    FILES_DIR,
    FULL_CONTRACT_PRICE_FIELD,
    FULL_EIS_CONTRACT_PRICE_FIELD,
    HEADER_ROW,
    STAGE_NAMES,
    STAGE_SUFFIXES,
    TARGET_BUDGET_YEAR,
)


@dataclass
class ProcessResult:
    status: str
    warning_count: int = 0
    changes: list[CellChange] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    control_changes: list[CellChange] = field(default_factory=list)
    price_report_lines: list[str] = field(default_factory=list)
    fas_status: str = ""
    fas_report_lines: list[str] = field(default_factory=list)
    publication_changed: bool = False


@dataclass
class RunMainResult:
    report_lines: list[str] = field(default_factory=list)
    changed_publication_notices: list[str] = field(default_factory=list)


@dataclass
class FasApplyResult:
    status: str = ""
    warnings: list[str] = field(default_factory=list)
    report_lines: list[str] = field(default_factory=list)


@dataclass
class NoticeTarget:
    sheet: object
    values_sheet: object
    row: int


DOMESTIC_SHARE_FIELD = "Доля ЛП отечественного производства, %"
FOREIGN_SHARE_FIELD = "Доля ЛП иностранного производства, %"
CONTRACT_EXECUTION_STATUS_FIELD = "Статус исполнения Контракта"
CONTRACT_EXECUTION_STATUS_VALUE = "в стадии исполнения"


def merge_data(eis_data, protocol_data, contract_data):
    result = {}
    result.update(contract_data.get("technical_characteristics", {}))
    result.update(contract_data.get("delivery_terms", {}))
    result.update(contract_data.get("funding_terms", {}))
    result.update(contract_data.get("contract_terms", {}))
    result.update(protocol_data.get("ppiop", {}))

    eis_fields = dict(eis_data.get("eis", {}))
    eis_contract_price = eis_fields.pop("Цена контракта", None)
    result.update(eis_fields)
    apply_eis_price_control(result, eis_contract_price)
    preserve_full_contract_price(result)
    apply_contract_year_metadata(result)
    apply_country_origin_shares(result)
    result[CONTRACT_EXECUTION_STATUS_FIELD] = CONTRACT_EXECUTION_STATUS_VALUE

    return result


def apply_contract_year_metadata(result):
    budget_years = normalize_years(result.get("Годы финансирования"))
    delivery_years = get_delivery_years(result)

    if not delivery_years:
        delivery_years = budget_years

    if budget_years:
        result["Бюджет"] = format_year_period(budget_years)
    elif delivery_years:
        result["Бюджет"] = format_year_period(delivery_years)
    else:
        result["Бюджет"] = str(TARGET_BUDGET_YEAR)

    if delivery_years:
        result["Период поставок по годам"] = format_year_period(delivery_years)


def normalize_years(value):
    if value in (None, ""):
        return []

    if isinstance(value, (list, tuple, set)):
        years = []
        for item in value:
            years.extend(normalize_years(item))
        return sorted(set(years))

    return sorted({int(year) for year in re.findall(r"\b20\d{2}\b", str(value))})


def get_delivery_years(data):
    years = sorted({
        stage["year"]
        for stage in collect_delivery_stages(data)
        if stage["year"] is not None
    })
    if years:
        return years

    return normalize_years(data.get("Период поставок по годам"))


def format_year_period(years):
    years = sorted(set(int(year) for year in years if year))
    if not years:
        return ""

    return "-".join(str(year) for year in years)


def preserve_full_contract_price(result):
    contract_price = result.get("Цена контракта")
    if is_numeric_price(contract_price):
        result[FULL_CONTRACT_PRICE_FIELD] = contract_price


def collect_delivery_stages(result):
    stages = []

    for stage in STAGE_NAMES:
        suffix = STAGE_SUFFIXES[stage]
        quantity_field = f"Количество товара по {suffix}"
        delivery_field = f"Срок поставки по {suffix}"
        document_field = f"Срок предоставления документов по {suffix}"
        quantity = get_number_value(result, quantity_field)
        delivery_date = get_date_value(result, delivery_field)
        document_date = get_date_value(result, document_field)

        if quantity is None and delivery_date is None and document_date is None:
            continue

        stage_year = None
        if delivery_date:
            stage_year = delivery_date.year
        elif document_date:
            stage_year = document_date.year

        stages.append({
            "stage": stage,
            "quantity": quantity,
            "quantity_field": quantity_field,
            "delivery_date": delivery_date,
            "document_date": document_date,
            "year": stage_year,
        })

    return stages


def normalize_number_for_excel(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)

    return value


def apply_country_origin_shares(result):
    countries = extract_country_origin_items(result.get("Страна происхождения Товара"))
    if not countries:
        return

    total = len(countries)
    domestic_count = sum(1 for country in countries if is_domestic_country(country))
    domestic_share = round(domestic_count * 100 / total)
    foreign_share = 100 - domestic_share

    result[DOMESTIC_SHARE_FIELD] = domestic_share
    result[FOREIGN_SHARE_FIELD] = foreign_share


def extract_country_origin_items(value):
    text = str(value or "").strip()
    if not text:
        return []

    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", ";")
    text = re.sub(r"\b\d+\.\s*", ";", text)

    countries = []
    seen = set()
    for part in re.split(r";+", text):
        country = part.strip(" .;")
        if not country:
            continue

        key = normalize_country_key(country)
        if not key or key in seen:
            continue

        seen.add(key)
        countries.append(country)

    return countries


def normalize_country_key(value):
    value = str(value or "").lower().replace("ё", "е")
    return re.sub(r"[^а-яa-z]", "", value)


def is_domestic_country(country):
    key = normalize_country_key(country)
    return "российскаяфедерация" in key or key == "россия"


def apply_eis_price_control(result, eis_contract_price):
    if not is_numeric_price(eis_contract_price):
        return

    result[FULL_EIS_CONTRACT_PRICE_FIELD] = eis_contract_price
    word_contract_price = result.get("Цена контракта")
    if not is_numeric_price(word_contract_price):
        result["Цена контракта ЕИС"] = eis_contract_price
        return

    if values_are_equal(word_contract_price, eis_contract_price, "Цена контракта"):
        result["Цена контракта ЕИС"] = ""
    else:
        result["Цена контракта ЕИС"] = eis_contract_price


def is_numeric_price(value):
    return get_number_value({"Цена контракта": value}, "Цена контракта") is not None


def apply_fas_data(data, notice_number):
    price = get_number_value(data, "Цена контракта")
    if price is None or price <= FAS_PRICE_THRESHOLD:
        return FasApplyResult()

    admitted_applications = get_number_value(data, "Кол-во допущенных заявок")
    if admitted_applications is not None and admitted_applications > 1:
        message = (
            f'цена контракта больше {format_number_for_message(FAS_PRICE_THRESHOLD)}, '
            f'допущенных заявок {format_number_for_message(admitted_applications)}'
        )
        return FasApplyResult(
            status="ФАС не требуется",
            report_lines=[f'    - ФАС не требуется: {message}'],
        )

    decision = get_fas_decision(notice_number)
    if not decision:
        warning = (
            f'цена контракта больше {format_number_for_message(FAS_PRICE_THRESHOLD)}, '
            'но решение ФАС не найдено'
        )
        return FasApplyResult(
            status="ФАС не найден",
            warnings=[warning],
            report_lines=[f'    - ФАС не найден: {warning}'],
        )

    data[FAS_APPROVAL_FIELD] = decision.value
    data[FAS_NOTIFICATION_FIELD] = decision.value

    validation = validate_fas_decision(decision, data)
    markers = ', '.join(validation.matched_markers) if validation.matched_markers else 'только номер извещения'
    report_lines = [
        f'    - ФАС заполнен: {decision.value}; файл {decision.source_name}; проверка: {markers}'
    ]
    report_lines.extend(f'    - ФАС проверка: {warning}' for warning in validation.warnings)

    return FasApplyResult(
        status="ФАС заполнен",
        warnings=list(validation.warnings),
        report_lines=report_lines,
    )


def format_number_for_message(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)

    if isinstance(value, int):
        return f'{value:,}'.replace(',', ' ')

    return str(value)


def build_notice_targets(workbook, values_workbook):
    targets = {}
    for sheet in get_publication_sheets(workbook):
        values_sheet = values_workbook[sheet.title]
        for row, notice_number in iter_notice_rows(sheet):
            targets.setdefault(notice_number, []).append(NoticeTarget(
                sheet=sheet,
                values_sheet=values_sheet,
                row=row,
            ))

    return targets


def get_sheet_year(sheet):
    match = re.search(r"\b(20\d{2})\b", str(sheet.title))
    if match:
        return int(match.group(1))

    return None


def select_publication_targets(data, targets):
    warnings = []
    delivery_years = get_delivery_years(data)
    if not delivery_years:
        delivery_years = sorted({
            year
            for year in (get_sheet_year(target.sheet) for target in targets)
            if year is not None
        })

    if not delivery_years and targets:
        target = targets[0]
        return [(target, dict(data))], warnings

    selected = []
    for delivery_year in delivery_years:
        year_targets = [
            target
            for target in targets
            if get_sheet_year(target.sheet) == delivery_year
        ]
        if not year_targets:
            warnings.append(f"нет строки на листе {delivery_year} год")
            continue

        year_data = build_year_publication_data(data, delivery_year)
        selected.extend((target, year_data) for target in year_targets)

    return selected, warnings


def build_year_publication_data(data, target_year):
    year_data = dict(data)
    stages = collect_delivery_stages(data)
    if not stages:
        return year_data

    selected_stages = [
        stage
        for stage in stages
        if stage["year"] == target_year or stage["year"] is None
    ]
    if not selected_stages:
        return year_data

    clear_stage_fields(year_data)

    original_total_quantity = sum(
        stage["quantity"]
        for stage in stages
        if stage["quantity"] is not None
    )
    target_total_quantity = sum(
        stage["quantity"]
        for stage in selected_stages
        if stage["quantity"] is not None
    )

    for output_stage, source_stage in zip(STAGE_NAMES, selected_stages):
        output_suffix = STAGE_SUFFIXES[output_stage]
        year_data[f"Количество товара по {output_suffix}"] = normalize_number_for_excel(
            source_stage["quantity"]
        )
        year_data[f"Срок поставки по {output_suffix}"] = source_stage["delivery_date"]
        year_data[f"Срок предоставления документов по {output_suffix}"] = source_stage["document_date"]

    if target_total_quantity:
        year_data["Общее количество товара"] = normalize_number_for_excel(target_total_quantity)
        apply_year_price_scope(year_data, data, original_total_quantity, target_total_quantity)

    return year_data


def clear_stage_fields(data):
    for stage in STAGE_NAMES:
        suffix = STAGE_SUFFIXES[stage]
        data[f"Количество товара по {suffix}"] = ""
        data[f"Срок поставки по {suffix}"] = ""
        data[f"Срок предоставления документов по {suffix}"] = ""


def apply_year_price_scope(year_data, source_data, original_total_quantity, target_total_quantity):
    if not original_total_quantity or original_total_quantity == target_total_quantity:
        return

    for field_key in ("Цена контракта", "Цена контракта ЕИС"):
        price = get_number_value(source_data, field_key)
        if price is None:
            continue

        unit_price = price / original_total_quantity
        year_data[field_key] = normalize_number_for_excel(
            round(unit_price * target_total_quantity, 2)
        )


def process_notice(
    notice_number,
    targets,
    journal_sheet,
    files_dir=FILES_DIR,
    journal_values_sheet=None,
    execution_sheet=None,
    execution_values_sheet=None,
    write_enabled=True,
):
    try:
        document_pair = find_document_pair(notice_number, files_dir)
    except ValueError as e:
        print(f'ERR {notice_number}: ошибка документов, {e}')
        return ProcessResult("document_error")

    if document_pair is None:
        return ProcessResult("skipped")

    protocol_file, contract_file = document_pair

    try:
        eis_data = fetch_eis_data(notice_number)
    except Exception as e:
        print(f'ERR {notice_number}: ошибка ЕИС, {format_eis_error(e)}')
        return ProcessResult("eis_error")

    try:
        protocol_data, contract_data = parse_document_pair(
            str(protocol_file),
            str(contract_file)
        )
    except Exception as e:
        print(f'ERR {notice_number}: ошибка документов ({e})')
        return ProcessResult("document_error")

    data = merge_data(eis_data, protocol_data, contract_data)
    fas_result = apply_fas_data(data, notice_number)
    filled_fields, total_fields = count_excel_fields(data)

    if not validate_data(data):
        missing_fields = ', '.join(get_missing_fields(data))
        print(
            f'ERR {notice_number}: ошибка проверки данных, '
            f'данные Excel {filled_fields}/{total_fields}, не найдены поля: {missing_fields}'
        )
        return ProcessResult("validation_error")

    suspicious_warnings = fas_result.warnings + check_suspicious_data(data)
    for warning in suspicious_warnings:
        print(f'WARN {notice_number}: {warning}')

    publication_targets, target_warnings = select_publication_targets(data, targets)
    if not publication_targets:
        return ProcessResult(
            "validation_error",
            warning_count=len(suspicious_warnings) + len(target_warnings),
            warnings=suspicious_warnings + target_warnings,
        )

    publication_results = []
    for target, year_data in publication_targets:
        publication_results.append(write_data_to_row(
            target.sheet,
            year_data,
            target.row,
            notice_number,
            target.values_sheet,
            apply_changes=write_enabled,
            sheet_name=target.sheet.title,
        ))

    primary_target = publication_targets[0][0]
    journal_result = write_journal_register(
        public_sheet=primary_target.sheet,
        journal_sheet=journal_sheet,
        row=primary_target.row,
        data=data,
        notice_number=notice_number,
        journal_values_sheet=journal_values_sheet,
        public_values_sheet=primary_target.values_sheet,
        apply_changes=write_enabled,
    )
    execution_result = write_execution_by_month(
        execution_sheet=execution_sheet,
        public_sheet=primary_target.sheet,
        row=primary_target.row,
        data=data,
        public_values_sheet=primary_target.values_sheet,
        execution_values_sheet=execution_values_sheet,
        apply_changes=write_enabled,
    )

    publication_warnings = target_warnings[:]
    publication_changes = []
    publication_changed = False
    for publication_result in publication_results:
        publication_warnings.extend(publication_result.warnings)
        publication_changes.extend(publication_result.changes)
        publication_changed = publication_result.changed or publication_changed

    for warning in publication_warnings:
        print(f'WARN {notice_number}: {warning}')
    for warning in journal_result.warnings:
        print(f'WARN {notice_number}: {warning}')
    for warning in execution_result.warnings:
        print(f'WARN {notice_number}: {warning}')

    changed = publication_changed or journal_result.changed or execution_result.changed
    publication_status, journal_status, execution_status = format_sheet_statuses(
        publication_changed,
        journal_result.changed,
        execution_result.changed,
        write_enabled,
    )
    warnings = (
        suspicious_warnings
        + publication_warnings
        + journal_result.warnings
        + execution_result.warnings
    )
    changes = publication_changes + journal_result.changes + execution_result.changes
    control_changes = [
        change
        for change in publication_changes
        if is_delivery_control_change(change) and change.kind != 'FORMAT'
    ]
    price_report_lines = build_price_report_lines(data)
    status_prefix = 'CHANGE' if control_changes or price_report_lines else 'OK'
    fas_status_suffix = f'; {fas_result.status}' if fas_result.status else ''

    if changed:
        print(
            f'{status_prefix}  {notice_number}: найдено 2/2, данные Excel {filled_fields}/{total_fields}; '
            f'{publication_status}; {journal_status}; {execution_status}{fas_status_suffix}'
        )
        return ProcessResult(
            status="processed",
            warning_count=len(warnings),
            changes=changes,
            warnings=warnings,
            control_changes=control_changes,
            price_report_lines=price_report_lines,
            fas_status=fas_result.status,
            fas_report_lines=fas_result.report_lines,
            publication_changed=publication_changed,
        )

    print(
        f'{status_prefix}  {notice_number}: найдено 2/2, данные Excel {filled_fields}/{total_fields}; '
        f'{publication_status}; {journal_status}; {execution_status}{fas_status_suffix}'
    )
    return ProcessResult(
        status="unchanged",
        warning_count=len(warnings),
        changes=changes,
        warnings=warnings,
        control_changes=control_changes,
        price_report_lines=price_report_lines,
        fas_status=fas_result.status,
        fas_report_lines=fas_result.report_lines,
        publication_changed=publication_changed,
    )


def format_sheet_statuses(publication_changed, journal_changed, execution_changed, write_enabled=True):
    if write_enabled:
        publication_status = 'Публикация изменено' if publication_changed else 'Публикация без изменений'
        journal_status = 'Журнал регистр внесено' if journal_changed else 'Журнал регистр без изменений'
        execution_status = 'Исп по мес внесено' if execution_changed else 'Исп по мес без изменений'
    else:
        publication_status = 'Публикация будет изменено' if publication_changed else 'Публикация без изменений'
        journal_status = 'Журнал регистр будет внесено' if journal_changed else 'Журнал регистр без изменений'
        execution_status = 'Исп по мес будет внесено' if execution_changed else 'Исп по мес без изменений'

    return publication_status, journal_status, execution_status


def clear_sheet_values_below_header(sheet):
    cleared_cells = 0

    for row in sheet.iter_rows(
        min_row=HEADER_ROW + 1,
        max_row=sheet.max_row,
        max_col=sheet.max_column,
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.value in (None, ""):
                continue

            cell.value = None
            cleared_cells += 1

    return cleared_cells


def clear_output_sheets(workbook, values_workbook):
    journal_sheet = workbook[JOURNAL_SHEET_NAME]
    journal_values_sheet = values_workbook[JOURNAL_SHEET_NAME]
    execution_sheet = find_execution_sheet(workbook)
    execution_values_sheet = find_execution_sheet(values_workbook)

    journal_cells = clear_sheet_values_below_header(journal_sheet)
    execution_cells = clear_sheet_values_below_header(execution_sheet)

    clear_sheet_values_below_header(journal_values_sheet)
    clear_sheet_values_below_header(execution_values_sheet)

    return journal_cells, execution_cells


def run_main(workbook_path=EXCEL_PATH, files_dir=FILES_DIR, write_enabled=True):
    if write_enabled:
        try:
            ensure_excel_is_closed(workbook_path)
        except RuntimeError as e:
            print(f'ERR: {e}')
            return RunMainResult()
    else:
        print('Режим проверки: запись в Excel отключена')

    wb = load_workbook(workbook_path)
    values_wb = load_workbook(workbook_path, data_only=True)
    structure_errors = validate_workbook_structure(wb)
    if structure_errors:
        print('ERR: структура Excel не подходит')
        for error in structure_errors:
            print(f'ERR: {error}')
        return RunMainResult()

    if write_enabled:
        journal_cleared, execution_cleared = clear_output_sheets(wb, values_wb)
        print(
            'Очищены листы перед запуском: '
            f'{JOURNAL_SHEET_NAME} ({journal_cleared} ячеек), '
            f'{find_execution_sheet(wb).title} ({execution_cleared} ячеек)'
        )

    journal_sheet = wb[JOURNAL_SHEET_NAME]
    journal_values_sheet = values_wb[JOURNAL_SHEET_NAME]
    execution_sheet = find_execution_sheet(wb)
    execution_values_sheet = find_execution_sheet(values_wb)
    notice_targets = build_notice_targets(wb, values_wb)

    processed = 0
    unchanged = 0
    skipped = 0
    eis_errors = 0
    document_errors = 0
    validation_errors = 0
    suspicious_warnings = 0
    report_lines = []
    changed_publication_notices = []

    for notice_number, targets in sorted(
        notice_targets.items(),
        key=lambda item: (
            get_sheet_year(item[1][0].sheet) or 9999,
            item[1][0].row,
            item[0],
        ),
    ):
        try:
            result = process_notice(
                notice_number,
                targets,
                journal_sheet,
                files_dir,
                journal_values_sheet,
                execution_sheet,
                execution_values_sheet,
                write_enabled,
            )
            status = result.status
            suspicious_warnings += result.warning_count
            report_lines.extend(build_notice_report_lines(notice_number, result))
            if result.publication_changed:
                changed_publication_notices.append(notice_number)
        except Exception as e:
            validation_errors += 1
            print(f'ERR {notice_number}: непредвиденная ошибка ({e})')
            continue

        if status == "processed":
            processed += 1
        elif status == "unchanged":
            unchanged += 1
        elif status == "skipped":
            skipped += 1
        elif status == "eis_error":
            eis_errors += 1
        elif status == "document_error":
            document_errors += 1
        else:
            validation_errors += 1

    if processed and write_enabled:
        prepare_workbook_for_recalculation(wb)
        wb.save(workbook_path)

    print('')
    print('Итог:')
    if write_enabled:
        print(f'Записано в Excel: {processed}')
    else:
        print(f'Будет изменено в Excel: {processed}')
    print(f'Без изменений: {unchanged}')
    print(f'Пропущено без документов: {skipped}')
    print(f'Ошибок ЕИС: {eis_errors}')
    print(f'Ошибок документов: {document_errors}')
    print(f'Ошибок проверки данных: {validation_errors}')
    print(f'Подозрительных предупреждений: {suspicious_warnings}')

    return RunMainResult(
        report_lines=report_lines,
        changed_publication_notices=changed_publication_notices,
    )


def prepare_workbook_for_recalculation(workbook):
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def build_notice_report_lines(notice_number, result):
    if (
        not result.control_changes
        and not result.price_report_lines
        and not result.fas_report_lines
    ):
        return []

    lines = [f'{notice_number}:']
    for change in result.control_changes:
        lines.append(format_change_line(change))
    lines.extend(result.price_report_lines)
    lines.extend(result.fas_report_lines)

    return lines


def build_price_report_lines(data):
    word_price = data.get("Цена контракта")
    eis_price = data.get("Цена контракта ЕИС")

    if not is_numeric_price(word_price) or not is_numeric_price(eis_price):
        return []

    return [
        '    - Изменение цены контракта: '
        f'Word/AE {format_report_value(word_price)}, '
        f'ЕИС/AF {format_report_value(eis_price)}'
    ]


def is_delivery_control_change(change):
    field = change.field.replace(' (формат)', '')
    return (
        field == 'Общее количество товара'
        or field.startswith('Количество товара по ')
        or field.startswith('Срок поставки по ')
        or field.startswith('Срок предоставления документов по ')
    )


def format_change_line(change):
    return (
        f'    - {change.sheet_name} {change.cell} {change.field}: '
        f'было {format_report_value(change.old_value)}, '
        f'стало {format_report_value(change.new_value)}'
    )


def format_report_value(value):
    if value in (None, ''):
        return 'пусто'
    if isinstance(value, datetime):
        return value.strftime('%d.%m.%Y')
    if isinstance(value, date):
        return value.strftime('%d.%m.%Y')
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value).is_integer():
            return str(int(value))
        return f'{float(value):.2f}'.replace('.', ',')

    return str(value).replace('\r\n', '\n').replace('\r', '\n').replace('\n', ' / ')


def write_detailed_report(log_file, report_lines):
    log_file.write('\n\nОтчет изменений по контролю поставки, цене и ФАС:\n')
    if not report_lines:
        log_file.write('Изменений по контролируемым полям, цене и ФАС нет\n')
    else:
        for line in report_lines:
            log_file.write(f'{line}\n')
    log_file.flush()


def main(
    workbook_path=EXCEL_PATH,
    files_dir=FILES_DIR,
    write_enabled=True,
    refresh_publication_from_main=False,
    apply_changed_filter=False,
):
    log_path = create_log_path()
    original_stdout = sys.stdout
    run_result = RunMainResult()

    with open(log_path, 'w', encoding='utf-8') as log_file:
        sys.stdout = TeeOutput(original_stdout, log_file)
        try:
            if refresh_publication_from_main and write_enabled:
                import_result = copy_main_publication_to_shapka(target_path=workbook_path)
                if import_result.errors:
                    for error in import_result.errors:
                        print(f'ERR: {error}')
                    write_detailed_report(log_file, [])
                    print(f'Лог сохранен: {log_path.relative_to(BASE_DIR)}')
                    return run_result

                print(
                    'Публикация обновлена из основной таблицы: '
                    f'{import_result.copied_rows} строк, {import_result.copied_columns} колонок'
                )

            run_result = run_main(workbook_path, files_dir, write_enabled=write_enabled)

            if apply_changed_filter and write_enabled:
                filter_result = apply_publication_changed_filter(
                    run_result.changed_publication_notices,
                    workbook_path=workbook_path,
                )
                if filter_result.errors:
                    for error in filter_result.errors:
                        print(f'ERR: {error}')
                elif filter_result.notice_count:
                    print(f'Фильтр включен по измененным строкам: {filter_result.notice_count}')
                else:
                    print('Фильтр по измененным строкам не включен: измененных строк нет')

            write_detailed_report(log_file, run_result.report_lines)
            print(f'Лог сохранен: {log_path.relative_to(BASE_DIR)}')
        finally:
            sys.stdout = original_stdout

    return run_result


def parse_args():
    parser = argparse.ArgumentParser(description='Парсер контрактов и протоколов')
    parser.add_argument(
        '--write',
        action='store_true',
        required=True,
        help='записать данные в Excel',
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    main(write_enabled=args.write)
