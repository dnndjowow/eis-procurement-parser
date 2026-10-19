import re

from openpyxl.utils import get_column_letter

from config import field_mappings_public, optional_field_keys
from settings import HEADER_ROW, NOTICE_COLUMN


PUBLICATION_SHEET_NAME = "Публикация"
JOURNAL_SHEET_NAME = "Журнал регистр"
EXECUTION_SHEET_ALIASES = {"исппомес"}
PUBLICATION_YEAR_SHEET_PATTERN = re.compile(r"^20\d{2}\s+год$")

JOURNAL_FIELD_MAPPINGS = [
    ("Номер извещения", "основание для заключения"),
    ("Дата заключения контракта", "Дата контракта"),
    ("Номер контракта", "Номер контракта"),
    ("МНН закупаемого лекарственного препарата", "предмет гк/дс"),
    ("Поставщик", "наименование поставщика/исполнителя/подрядчика"),
    ("ИНН", "ИНН"),
    ("Цена контракта", "Цена контракта"),
    ("КОСГУ", "КОСГУ"),
    (
        "Нормативно-правовой акт, в рамках которого осуществляется централизованная закупка лекарственных препаратов",
        "Примечание",
    ),
]

EXECUTION_COLUMN_MAP = {
    "Нормативно-правовой акт": "A",
    "Поставщик": "B",
    "Дата контракта": "C",
    "Номер контракта": "D",
    "МНН": "E",
    "Сумма ГК/этапа к оплате, руб.": "F",
    "Кол-во товара": "G",
    "Цена за ед. товара, руб.": "H",
    "Срок поставки Товара": "I",
    "Срок предоставления документов в ЕИС": "J",
}

EXECUTION_FIELD_ALIASES = {
    "Нормативно-правовой акт": (
        "Нормативно-правовой акт, в рамках которого осуществляется централизованная закупка лекарственных препаратов",
    ),
    "Поставщик": ("Поставщик",),
    "Дата контракта": ("Дата контракта", "Дата ГК"),
    "Номер контракта": ("Номер контракта",),
    "МНН": ("МНН",),
    "Сумма ГК/этапа к оплате, руб.": ("Сумма ГК/этапа к оплате, руб.",),
    "Кол-во товара": ("кол-ва товара",),
    "Цена за ед. товара, руб.": ("цена за ед. товара, руб.",),
    "Срок поставки Товара": ("Срок поставки Товара, до", "Срок поставки Товара по ГК"),
    "Срок предоставления документов в ЕИС": (
        "Срок предоставления документов в ЕИС",
        "Срок предоставления документов в ЕИС, до",
    ),
}


def normalize_layout_header(value):
    return " ".join(str(value or "").replace("\n", " ").split()).lower()


def normalize_sheet_name(value):
    text = str(value or "").lower().replace("ё", "е")
    return re.sub(r"[^а-яa-z0-9]", "", text)


def is_publication_year_sheet(sheet_name):
    return bool(PUBLICATION_YEAR_SHEET_PATTERN.fullmatch(str(sheet_name or "").strip()))


def get_publication_sheets(workbook):
    return [
        workbook[sheet_name]
        for sheet_name in workbook.sheetnames
        if is_publication_year_sheet(sheet_name)
    ]


def find_execution_sheet(workbook):
    for sheet_name in workbook.sheetnames:
        if normalize_sheet_name(sheet_name) in EXECUTION_SHEET_ALIASES:
            return workbook[sheet_name]

    raise KeyError("Лист исп_по_мес не найден")


def build_layout_header_map(sheet, header_row=HEADER_ROW):
    header_columns = {}

    for cell in sheet[header_row]:
        if cell.value not in (None, ""):
            header_columns[normalize_layout_header(cell.value)] = get_column_letter(cell.column)

    return header_columns


def header_exists(sheet, column_or_header, header_columns=None):
    column_or_header = str(column_or_header).strip()
    if re.fullmatch(r"[A-Z]{1,3}", column_or_header):
        return True

    if header_columns is None:
        header_columns = build_layout_header_map(sheet)

    return normalize_layout_header(column_or_header) in header_columns


def header_matches(value, aliases):
    normalized_value = normalize_layout_header(value)
    return any(normalized_value == normalize_layout_header(alias) for alias in aliases)


def validate_publication_sheet(sheet, sheet_name=None):
    errors = []
    header_columns = build_layout_header_map(sheet)
    checked_headers = set()
    sheet_name = sheet_name or sheet.title

    if not header_exists(sheet, NOTICE_COLUMN, header_columns):
        errors.append(f"{sheet_name}: не найдена колонка с номером извещения")

    for field_key, column_or_header in field_mappings_public:
        if field_key in optional_field_keys:
            continue
        if column_or_header in checked_headers:
            continue
        checked_headers.add(column_or_header)

        if not header_exists(sheet, column_or_header, header_columns):
            errors.append(f"{sheet_name}: не найдена колонка: {column_or_header}")

    return errors


def validate_journal_sheet(sheet):
    errors = []
    header_columns = build_layout_header_map(sheet)

    for _, column_or_header in JOURNAL_FIELD_MAPPINGS:
        if not header_exists(sheet, column_or_header, header_columns):
            errors.append(f"{JOURNAL_SHEET_NAME}: не найдена колонка: {column_or_header}")

    return errors


def validate_execution_sheet(sheet):
    errors = []

    for field, column in EXECUTION_COLUMN_MAP.items():
        header_value = sheet[f"{column}{HEADER_ROW}"].value
        aliases = EXECUTION_FIELD_ALIASES[field]
        if not header_matches(header_value, aliases):
            errors.append(f"исп_по_мес: колонка {column} должна быть: {field}")

    if errors:
        errors.insert(0, "исп_по_мес: структура листа не соответствует Shapka_Example")

    return errors


def validate_workbook_structure(workbook):
    errors = []
    publication_sheets = get_publication_sheets(workbook)

    if not publication_sheets:
        errors.append("не найдены листы публикации по годам")
    else:
        for sheet in publication_sheets:
            errors.extend(validate_publication_sheet(sheet, sheet.title))

    if JOURNAL_SHEET_NAME not in workbook.sheetnames:
        errors.append(f"не найден лист {JOURNAL_SHEET_NAME}")
    else:
        errors.extend(validate_journal_sheet(workbook[JOURNAL_SHEET_NAME]))

    try:
        execution_sheet = find_execution_sheet(workbook)
    except KeyError as e:
        errors.append(str(e).strip("'"))
    else:
        errors.extend(validate_execution_sheet(execution_sheet))

    return errors
