import copy
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.filters import AutoFilter

from excel_tools import (
    ensure_excel_is_closed,
    extract_notice_number,
)
from excel_layout import (
    PUBLICATION_SHEET_NAME,
    get_publication_sheets,
    is_publication_year_sheet,
)
from settings import (
    EXCEL_PATH,
    MAIN_PUBLICATION_EXCEL_PATH,
    PUBLICATION_SYNC_LAST_COLUMN,
)


NOTICE_COLUMN_INDEX = 4
HEADER_ROW = 1
CHANGED_ROW_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")


@dataclass
class PublicationImportResult:
    errors: list[str] = field(default_factory=list)
    copied_rows: int = 0
    copied_columns: int = 0


@dataclass
class PublicationFilterResult:
    errors: list[str] = field(default_factory=list)
    notice_count: int = 0
    filled_rows: int = 0


def copy_main_publication_to_shapka(
    source_path=MAIN_PUBLICATION_EXCEL_PATH,
    target_path=EXCEL_PATH,
    target_sheet_name=PUBLICATION_SHEET_NAME,
):
    source_path = Path(source_path)
    target_path = Path(target_path)
    result = PublicationImportResult()

    try:
        ensure_excel_is_closed(source_path)
        ensure_excel_is_closed(target_path)
    except RuntimeError as e:
        result.errors.append(str(e))
        return result

    if not source_path.exists():
        result.errors.append(f"не найден файл основной таблицы: {source_path}")
        return result
    if not target_path.exists():
        result.errors.append(f"не найден файл Shapka_New: {target_path}")
        return result

    source_wb = load_workbook(source_path)
    target_wb = load_workbook(target_path)

    source_sheets = [
        sheet
        for sheet in source_wb.worksheets
        if is_publication_year_sheet(sheet.title)
    ]
    if not source_sheets:
        result.errors.append(f"не найдены годовые листы публикации в {source_path.name}")
        return result

    remove_target_publication_sheets(target_wb, target_sheet_name)
    for index, source_sheet in enumerate(source_sheets):
        copy_sheet_between_workbooks(source_sheet, target_wb, source_sheet.title, target_index=index)

    target_wb.save(target_path)

    result.copied_rows = sum(sheet.max_row for sheet in source_sheets)
    result.copied_columns = max(sheet.max_column for sheet in source_sheets)
    return result


def remove_target_publication_sheets(target_workbook, old_publication_sheet_name):
    for sheet_name in list(target_workbook.sheetnames):
        if sheet_name == old_publication_sheet_name or is_publication_year_sheet(sheet_name):
            del target_workbook[sheet_name]


def copy_sheet_between_workbooks(source_sheet, target_workbook, target_sheet_name, target_index=None):
    if target_sheet_name in target_workbook.sheetnames:
        if target_index is None:
            target_index = target_workbook.sheetnames.index(target_sheet_name)
        del target_workbook[target_sheet_name]
    elif target_index is None:
        target_index = len(target_workbook.sheetnames)

    target_sheet = target_workbook.create_sheet(target_sheet_name, target_index)
    target_sheet.sheet_state = source_sheet.sheet_state

    copy_sheet_properties(source_sheet, target_sheet)
    copy_dimensions(source_sheet, target_sheet)
    copy_cells(source_sheet, target_sheet)
    copy_merged_cells(source_sheet, target_sheet)


def copy_sheet_properties(source_sheet, target_sheet):
    target_sheet.sheet_format = copy.copy(source_sheet.sheet_format)
    target_sheet.sheet_properties = copy.copy(source_sheet.sheet_properties)
    target_sheet.page_margins = copy.copy(source_sheet.page_margins)
    target_sheet.page_setup = copy.copy(source_sheet.page_setup)
    target_sheet.print_options = copy.copy(source_sheet.print_options)
    if hasattr(source_sheet, "header_footer") and hasattr(target_sheet, "header_footer"):
        target_sheet.header_footer = copy.copy(source_sheet.header_footer)
    target_sheet.freeze_panes = source_sheet.freeze_panes
    target_sheet.auto_filter = copy.copy(source_sheet.auto_filter)

    if source_sheet.sheet_view:
        copy_sheet_view(source_sheet, target_sheet)

    target_sheet.print_area = source_sheet.print_area
    target_sheet.sheet_properties.tabColor = source_sheet.sheet_properties.tabColor


def copy_sheet_view(source_sheet, target_sheet):
    for attr in (
        "showGridLines",
        "showRowColHeaders",
        "showZeros",
        "rightToLeft",
        "tabSelected",
        "zoomScale",
        "zoomScaleNormal",
        "view",
    ):
        try:
            setattr(target_sheet.sheet_view, attr, copy.copy(getattr(source_sheet.sheet_view, attr)))
        except AttributeError:
            continue


def copy_dimensions(source_sheet, target_sheet):
    for key, source_dimension in source_sheet.column_dimensions.items():
        target_dimension = target_sheet.column_dimensions[key]
        copy_dimension_attrs(
            source_dimension,
            target_dimension,
            ("width", "hidden", "outlineLevel", "collapsed", "bestFit", "style"),
        )

    for key, source_dimension in source_sheet.row_dimensions.items():
        target_dimension = target_sheet.row_dimensions[key]
        copy_dimension_attrs(
            source_dimension,
            target_dimension,
            ("height", "hidden", "outlineLevel", "collapsed", "style"),
        )


def copy_dimension_attrs(source_dimension, target_dimension, attrs):
    for attr in attrs:
        try:
            setattr(target_dimension, attr, copy.copy(getattr(source_dimension, attr)))
        except AttributeError:
            continue


def copy_cells(source_sheet, target_sheet):
    for row in source_sheet.iter_rows():
        for source_cell in row:
            target_cell = target_sheet.cell(
                row=source_cell.row,
                column=source_cell.column,
                value=source_cell.value,
            )

            if source_cell.has_style:
                target_cell._style = copy.copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.font:
                target_cell.font = copy.copy(source_cell.font)
            if source_cell.fill:
                target_cell.fill = copy.copy(source_cell.fill)
            if source_cell.border:
                target_cell.border = copy.copy(source_cell.border)
            if source_cell.alignment:
                target_cell.alignment = copy.copy(source_cell.alignment)
            if source_cell.protection:
                target_cell.protection = copy.copy(source_cell.protection)
            target_cell.hyperlink = copy.copy(source_cell.hyperlink)
            target_cell.comment = copy.copy(source_cell.comment)


def copy_merged_cells(source_sheet, target_sheet):
    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))


def apply_publication_changed_filter(
    notice_numbers,
    workbook_path=EXCEL_PATH,
    sheet_name=PUBLICATION_SHEET_NAME,
    last_column=PUBLICATION_SYNC_LAST_COLUMN,
):
    result = PublicationFilterResult(notice_count=len(set(notice_numbers)))
    workbook_path = Path(workbook_path)

    try:
        ensure_excel_is_closed(workbook_path)
    except RuntimeError as e:
        result.errors.append(str(e))
        return result

    wb = load_workbook(workbook_path)
    publication_sheets = get_publication_sheets(wb)
    if not publication_sheets:
        result.errors.append(f"не найдены годовые листы публикации в {workbook_path.name}")
        return result

    notice_numbers = sorted(set(notice_numbers))

    for sheet in publication_sheets:
        last_row = max(sheet.max_row, find_last_notice_row(sheet))
        if notice_numbers:
            apply_changed_notice_filter(sheet, notice_numbers, last_column, last_row)
            apply_changed_notice_row_visibility(sheet, notice_numbers, last_row)
            result.filled_rows += apply_changed_row_fill(sheet, notice_numbers, last_column, last_row)
        else:
            sheet.auto_filter = AutoFilter(ref=f"A1:{last_column}{last_row}")
            reset_row_visibility(sheet, last_row)

    wb.save(workbook_path)
    return result


def find_last_notice_row(sheet):
    last_row = HEADER_ROW

    for row in range(2, sheet.max_row + 1):
        if extract_notice_number(sheet.cell(row, NOTICE_COLUMN_INDEX).value):
            last_row = row

    return last_row


def apply_changed_notice_filter(sheet, notice_numbers, last_column, last_row):
    sheet.auto_filter = AutoFilter(ref=f"A1:{last_column}{last_row}")
    sheet.auto_filter.add_filter_column(
        NOTICE_COLUMN_INDEX - 1,
        sorted(set(notice_numbers)),
    )


def apply_changed_notice_row_visibility(sheet, notice_numbers, last_row):
    notice_numbers = set(notice_numbers)
    sheet.row_dimensions[HEADER_ROW].hidden = False

    for row in range(HEADER_ROW + 1, last_row + 1):
        notice_number = extract_notice_number(sheet.cell(row, NOTICE_COLUMN_INDEX).value)
        sheet.row_dimensions[row].hidden = notice_number not in notice_numbers


def apply_changed_row_fill(sheet, notice_numbers, last_column, last_row):
    notice_numbers = set(notice_numbers)
    last_column_index = column_index_from_string(last_column)
    filled_rows = 0

    for row in range(HEADER_ROW + 1, last_row + 1):
        notice_number = extract_notice_number(sheet.cell(row, NOTICE_COLUMN_INDEX).value)
        if notice_number not in notice_numbers:
            continue

        filled_rows += 1
        for column in range(1, last_column_index + 1):
            sheet.cell(row, column).fill = CHANGED_ROW_FILL

    return filled_rows


def reset_row_visibility(sheet, last_row):
    for row in range(HEADER_ROW, last_row + 1):
        sheet.row_dimensions[row].hidden = False
