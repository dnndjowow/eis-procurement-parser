import re
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from docx_parser.common import normalize_text
from settings import FAS_DIR, NOTICE_NUMBER_PATTERN


FAS_APPROVAL_FIELD = "Согласование ФАС"
FAS_NOTIFICATION_FIELD = "Уведомление ФАС"


@dataclass(frozen=True)
class FasDecision:
    source_name: str
    notice_numbers: tuple[str, ...]
    number: str
    date: str
    text: str

    @property
    def value(self):
        return f"№ {self.number} от {self.date}"


@dataclass(frozen=True)
class FasValidation:
    matched_markers: tuple[str, ...]
    warnings: tuple[str, ...]


def get_fas_decision(notice_number):
    return build_fas_decision_index().get(str(notice_number))


def get_fas_decision_value(notice_number):
    decision = get_fas_decision(notice_number)
    return decision.value if decision else None


@lru_cache(maxsize=1)
def build_fas_decision_index():
    decisions = {}

    for name, content in iter_fas_documents():
        decision = parse_fas_decision(name, content)
        if not decision:
            continue

        for notice_number in decision.notice_numbers:
            decisions[notice_number] = decision

    return decisions


def iter_fas_documents(fas_dir=FAS_DIR):
    fas_dir = Path(fas_dir)
    if not fas_dir.exists():
        return

    for path in sorted(fas_dir.iterdir()):
        if path.name.startswith("~$") or path.name.startswith("._"):
            continue

        if path.suffix.lower() == ".docx":
            yield path.name, path.read_bytes()
        elif path.suffix.lower() == ".zip":
            yield from iter_zip_documents(path)


def iter_zip_documents(path):
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith("/") or not name.lower().endswith(".docx"):
                continue
            yield name, archive.read(name)


def parse_fas_decision(name, content):
    text = extract_docx_text(content)
    if not text:
        return None

    notice_numbers = tuple(sorted(set(re.findall(NOTICE_NUMBER_PATTERN, text))))
    decision_number = extract_decision_number(text)
    decision_date = extract_decision_date(text)

    if not notice_numbers or not decision_number or not decision_date:
        return None

    return FasDecision(name, notice_numbers, decision_number, decision_date, text)


def validate_fas_decision(decision, data):
    matched_markers = []
    warnings = []

    text = normalize_text(decision.text).lower()
    text_digits = re.sub(r"\D", "", decision.text)
    text_words = set(re.findall(r"[а-яёa-z0-9]{4,}", text, flags=re.IGNORECASE))

    inn = normalize_digits(data.get("ИНН"))
    if inn and len(inn) in (10, 12) and inn in text_digits:
        matched_markers.append("ИНН")

    supplier_tokens = significant_tokens(data.get("Поставщик"))
    if supplier_tokens and tokens_match_text(supplier_tokens, text, text_words):
        matched_markers.append("поставщик")

    product_tokens = []
    for field in (
        "МНН закупаемого лекарственного препарата",
        "Торговое наименование",
    ):
        product_tokens.extend(significant_tokens(data.get(field)))
    if product_tokens and tokens_match_text(product_tokens, text, text_words):
        matched_markers.append("препарат")

    price_variants = money_digit_variants(data.get("Цена контракта"))
    if price_variants and any(variant in text_digits for variant in price_variants):
        matched_markers.append("цена")

    has_check_data = bool(inn or supplier_tokens or product_tokens or price_variants)
    if has_check_data and not matched_markers:
        warnings.append(
            f"решение ФАС {decision.value} найдено по номеру извещения, "
            "но не подтвердилось по ИНН/поставщику/препарату/цене"
        )

    return FasValidation(tuple(matched_markers), tuple(warnings))


def normalize_digits(value):
    return re.sub(r"\D", "", str(value or ""))


def significant_tokens(value):
    text = normalize_text(value).lower().replace("ё", "е")
    text = re.sub(r"[\"'«»”“()]", " ", text)
    tokens = re.findall(r"[а-яa-z0-9]{4,}", text, flags=re.IGNORECASE)
    stop_words = {
        "общество",
        "ограниченной",
        "ответственностью",
        "акционерное",
        "компания",
        "закрытое",
        "открытое",
        "публичное",
        "товарищество",
        "группа",
        "препарат",
        "лекарственный",
        "лекарственного",
        "таблетки",
        "капсулы",
        "раствор",
        "покрытые",
        "оболочкой",
        "пленочной",
    }

    return [
        token
        for token in tokens
        if token not in stop_words and not token.isdigit()
    ]


def tokens_match_text(tokens, text, text_words):
    if not tokens:
        return False

    matched = 0
    for token in tokens:
        if token in text:
            matched += 1
            continue

        if any(SequenceMatcher(None, token, word).ratio() >= 0.88 for word in text_words):
            matched += 1

    required = 1 if len(tokens) <= 2 else 2
    return matched >= required


def money_digit_variants(value):
    text = normalize_text(value)
    if not text:
        return []

    match = re.search(r"(\d[\d\s\u202f]*)(?:[,.](\d{1,2}))?", text)
    if not match:
        return []

    rubles = normalize_digits(match.group(1))
    kopecks = normalize_digits(match.group(2) or "")
    if not rubles:
        return []

    variants = [rubles]
    if kopecks:
        variants.insert(0, rubles + kopecks[:2].ljust(2, "0"))

    return [variant for variant in variants if len(variant) >= 6]


def extract_docx_text(content):
    doc = Document(BytesIO(content))
    lines = []

    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            text = " | ".join(normalize_text(cell.text) for cell in row.cells)
            if text.strip(" |"):
                lines.append(text)

    return "\n".join(lines)


def extract_decision_number(text):
    patterns = [
        r"дел[оу]\s*№\s*(\d+\s*/\s*\d+\s*/\s*\d+\s*/\s*\d+)",
        r"\b(26\s*/\s*44\s*/\s*93\s*/\s*\d+)\b",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return normalize_decision_number(match.group(1))

    return ""


def normalize_decision_number(value):
    parts = [part.strip() for part in value.split("/")]
    return "/".join(parts)


def extract_decision_date(text):
    match = re.search(r"\b(\d{2}\.\d{2}\.20\d{2})\s+Москва\b", text)
    if match:
        return match.group(1)

    dates = []
    for value in re.findall(r"\b\d{2}\.\d{2}\.20\d{2}\b", text):
        try:
            dates.append(datetime.strptime(value, "%d.%m.%Y").date())
        except ValueError:
            continue

    if not dates:
        return ""

    return max(dates).strftime("%d.%m.%Y")
