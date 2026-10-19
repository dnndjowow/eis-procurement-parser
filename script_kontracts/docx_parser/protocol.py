import os
import re
from typing import Dict, List

from docx import Document

from .common import clean_price_value, collect_document_lines, normalize_text, price_to_float


PRICE_PATTERN = r'(\d[\d\s]*,\d{2})'
PRICE_CONTEXT_MARKERS = [
    'цена контракта',
    'ценовое предложение',
    'предложение о цене',
    'предложением о цене',
    'цена, предложенная участником',
    'цена предложенная участником',
    'цена победителя',
]
BAD_PRICE_CONTEXT_MARKERS = [
    'нмцк',
    'начальная максимальная цена',
    'начальная (максимальная) цена',
    'обеспечение',
    'штраф',
    'пеня',
    'пени',
    'аванс',
]
GOOD_LOCAL_PRICE_MARKERS = [
    'ценовое предложение',
    'предложение о цене',
    'предложением о цене',
    'цена, предложенная участником',
    'цена предложенная участником',
    'цена победителя',
]


def extract_protocol_date(doc: Document, full_text: str) -> str:
    # В обычных ППИОП дата лежит в первой таблице, в системных протоколах - там же с подписью.
    if doc.tables:
        for row in doc.tables[0].rows:
            for cell in row.cells:
                date_match = re.search(r'\b(\d{2}\.\d{2}\.\d{4})\b', normalize_text(cell.text))
                if date_match:
                    return date_match.group(1)

    date_match = re.search(
        r'Дата подведения итогов[^:]*:\s*(\d{2}\.\d{2}\.\d{4})',
        full_text
    )
    if date_match:
        return date_match.group(1)

    return "-"


def extract_system_protocol_rows(doc: Document) -> List[List[str]]:
    for table in doc.tables:
        if not table.rows:
            continue

        header = [normalize_text(cell.text).lower() for cell in table.rows[0].cells]
        header_text = ' | '.join(header)

        if 'ценовое предложение' in header_text and 'решение о соответствии' in header_text:
            return [
                [normalize_text(cell.text) for cell in row.cells]
                for row in table.rows
            ]

    return []


def extract_allowed_applications_count(doc: Document, full_text: str) -> str:
    system_rows = extract_system_protocol_rows(doc)

    if system_rows:
        header = [cell.lower() for cell in system_rows[0]]
        decision_col = next(
            (
                index
                for index, cell in enumerate(header)
                if 'решение о соответствии' in cell
            ),
            -1
        )

        if decision_col != -1:
            count = 0
            for row in system_rows[1:]:
                if len(row) <= decision_col or not row[0].isdigit():
                    continue

                decision = row[decision_col].lower()
                if 'соответствует требованиям' in decision and 'не соответствует' not in decision:
                    count += 1

            if count:
                return str(count)

    full_text_lower = full_text.lower()
    if re.search(r'подан[аоы]?\s+только\s+одн[а-я]*\s+заявк', full_text_lower):
        return "1"

    if re.search(r'только\s+одн[а-я]*\s+заявк[а-я\s]+соответствует\s+требованиям', full_text_lower):
        return "1"

    if 'единственной заявки' in full_text_lower:
        return "1"

    return "-"


def extract_protocol_price(doc: Document, lines: List[str], full_text: str) -> str:
    system_rows = extract_system_protocol_rows(doc)

    if system_rows:
        header = [cell.lower() for cell in system_rows[0]]
        price_col = next(
            (
                index
                for index, cell in enumerate(header)
                if 'ценовое предложение' in cell
            ),
            -1
        )
        decision_col = next(
            (
                index
                for index, cell in enumerate(header)
                if 'решение о соответствии' in cell
            ),
            -1
        )
        prices = []

        if price_col != -1 and decision_col != -1:
            for row in system_rows[1:]:
                if len(row) <= max(price_col, decision_col) or not row[0].isdigit():
                    continue

                decision = row[decision_col].lower()
                if 'соответствует требованиям' in decision and 'не соответствует' not in decision:
                    price_match = re.search(PRICE_PATTERN, row[price_col])
                    if price_match:
                        prices.append(price_match.group(1))

        if prices:
            return clean_price_value(min(prices, key=price_to_float))

    for line in lines:
        line_lower = line.lower()
        if 'предложением о цене контракта' in line_lower or 'предложением о цене' in line_lower:
            prices = re.findall(PRICE_PATTERN + r'\s*(?:руб|рублей|$)', line)
            if prices:
                return clean_price_value(prices[-1])

    context_price = extract_price_from_context(full_text)
    if context_price:
        return context_price

    return extract_price_from_any_number(full_text)


def extract_price_from_context(full_text: str) -> str:
    candidates = []
    full_text_lower = full_text.lower()

    for marker in PRICE_CONTEXT_MARKERS:
        for marker_match in re.finditer(re.escape(marker), full_text_lower):
            start = max(0, marker_match.start() - 300)
            end = min(len(full_text), marker_match.end() + 300)
            context = full_text[start:end]
            marker_position = marker_match.start() - start

            for price_match in re.finditer(PRICE_PATTERN, context):
                candidate = build_price_candidate(
                    price_match.group(1),
                    context,
                    price_match.start(),
                    marker_position,
                )
                if candidate:
                    candidates.append(candidate)

    if not candidates:
        return ""

    clean_candidates = [candidate for candidate in candidates if not candidate["bad_context"]]
    if clean_candidates:
        candidates = clean_candidates

    return clean_price_value(min(candidates, key=lambda item: (item["distance"], item["value"]))["price"])


def build_price_candidate(price: str, context: str, price_position: int, marker_position: int):
    try:
        price_num = price_to_float(price)
    except ValueError:
        return None

    if price_num <= 0:
        return None

    return {
        "price": price,
        "value": price_num,
        "distance": abs(price_position - marker_position),
        "bad_context": has_bad_price_context(context, price_position),
    }


def has_bad_price_context(context: str, price_position: int) -> bool:
    start = max(0, price_position - 80)
    end = min(len(context), price_position + 80)
    local_context = context[start:end].lower()
    preceding_context = context[start:price_position].lower()
    if any(marker in preceding_context for marker in GOOD_LOCAL_PRICE_MARKERS):
        return False

    return any(marker in local_context for marker in BAD_PRICE_CONTEXT_MARKERS)


def extract_price_from_any_number(full_text: str) -> str:
    all_price_matches = re.findall(PRICE_PATTERN, full_text)
    prices_over_limit = []

    for price in all_price_matches:
        try:
            price_num = price_to_float(price)
        except ValueError:
            continue

        if price_num > 13000:
            prices_over_limit.append(price)

    if prices_over_limit:
        return clean_price_value(min(prices_over_limit, key=price_to_float))

    if all_price_matches:
        return clean_price_value(min(all_price_matches, key=price_to_float))

    return "-"


def parse_protocol(protocol_file: str) -> Dict[str, Dict[str, str]]:
    if not os.path.exists(protocol_file):
        raise FileNotFoundError(f"Файл протокола не найден: {protocol_file}")

    doc = Document(protocol_file)
    lines = collect_document_lines(doc)
    full_text = '\n'.join(lines)

    return {
        "ppiop": {
            "Дата итогового протокола": extract_protocol_date(doc, full_text),
            "Кол-во допущенных заявок": extract_allowed_applications_count(doc, full_text),
            "Цена по протоколу": extract_protocol_price(doc, lines, full_text),
        }
    }
