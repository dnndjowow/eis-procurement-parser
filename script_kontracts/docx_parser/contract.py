import os
import re
from typing import Any, Dict, List

from docx import Document

from .common import collect_document_lines, normalize_text
from .delivery import parse_delivery_terms


COUNTRY_NAMES = [
    "Российская Федерация",
    "Республика Беларусь",
    "Республика Казахстан",
    "Французская Республика",
    "Федеративная Республика Германия",
    "Соединенные Штаты Америки",
    "Республика Индия",
    "Королевство Испания",
    "Австрийская Республика",
    "Королевство Швеция",
    "Итальянская Республика",
    "Королевство Дания",
    "Швейцарская Конфедерация",
    "Королевство Нидерландов",
    "Королевство Бельгия",
    "Соединенное Королевство Великобритании и Северной Ирландии",
    "Венгрия",
    "Республика Словения",
    "Республика Корея",
    "Республика Пуэрто-Рико",
    "Ирландия",
    "Канада",
    "Япония",
    "Индия",
    "Китайская Народная Республика",
    "Китай",
]

COUNTRY_ALIASES = [
    ("Соединенные Штаты Америки", ["соединенные штаты", "сша", "united states"]),
    ("Российская Федерация", ["российская федерация", "россия"]),
    ("Федеративная Республика Германия", ["федеративная республика германия", "германия"]),
    ("Итальянская Республика", ["итальянская республика", "италия"]),
    ("Королевство Дания", ["королевство дания", "дания"]),
    ("Швейцарская Конфедерация", ["швейцарская конфедерация", "швейцария"]),
    ("Королевство Нидерландов", ["королевство нидерландов", "нидерланды"]),
    ("Соединенное Королевство Великобритании и Северной Ирландии", ["великобритания", "соединенное королевство", "united kingdom"]),
    ("Республика Корея", ["республика корея", "южная корея", "корея"]),
]


def clean_country_origin_value(value: str) -> str:
    value = normalize_text(value)
    value_lower = value.lower()

    found_countries = []
    for canonical_country, aliases in COUNTRY_ALIASES:
        if any(alias in value_lower for alias in aliases) and canonical_country not in found_countries:
            found_countries.append(canonical_country)

    for country in COUNTRY_NAMES:
        if country.lower() in value_lower and country not in found_countries:
            found_countries.append(country)

    for part in re.split(r';|\n|\b\d+\.\s*', value):
        part = part.strip(' .;')
        if not part:
            continue

        part_lower = part.lower()
        has_known_country = False
        for canonical_country, aliases in COUNTRY_ALIASES:
            if any(alias in part_lower for alias in aliases):
                has_known_country = True
                break
        if not has_known_country:
            has_known_country = any(country.lower() in part_lower for country in COUNTRY_NAMES)
        if has_known_country:
            continue

        if ',' in part:
            part = part.split(',', 1)[0].strip()
        if part and part not in found_countries:
            found_countries.append(part)

    return '; '.join(found_countries)


def is_valid_country_origin(value: str) -> bool:
    value = clean_country_origin_value(value)
    value_lower = value.lower()

    if not value or value == "-":
        return False
    if any(marker in value_lower for marker in [
        "страна",
        "товар",
        "единица",
        "количество",
        "сертификат",
        "подтверждение",
        "итого",
    ]):
        return False

    return True


def normalize_unit_value(value: str) -> str:
    value = normalize_text(value).rstrip("*").strip()
    replacements = {
        "шт.": "штука",
        "шт": "штука",
        "штук": "штука",
    }

    return replacements.get(value.lower(), value)


def normalize_form_value(value: str) -> str:
    value = normalize_text(value)
    return re.sub(r"(\d)\s*[xх]\s*(\d)", r"\1x\2", value, flags=re.IGNORECASE)


def is_registration_value(value: str) -> bool:
    value_lower = normalize_text(value).lower()

    if any(marker in value_lower for marker in [
        "держатель",
        "владелец",
        "производитель",
        "регистрационного удостоверения лекарственного препарата:",
    ]):
        return False

    return bool(re.search(r'\b(?:лп|лср|лс)\s*[-№]?\s*[\w()/.-]+|\bп\s*n\s*\d', value_lower))


def is_quantity_value(value: str) -> bool:
    value = normalize_text(value)

    if not value or value == "-":
        return False
    if re.search(r'[A-Za-zА-Яа-я]', value):
        return False

    return bool(re.search(r'\d', value))


def unique_values(values: List[str]) -> List[str]:
    result = []
    for value in values:
        value = normalize_text(value)
        if value and value != "-" and value not in result:
            result.append(value)
    return result


def format_values(values: List[str]) -> str:
    values = unique_values(values)

    if not values:
        return ""
    if len(values) == 1:
        return values[0]

    return "; ".join(f"{index}. {value}" for index, value in enumerate(values, 1))


def normalize_digits(value: str) -> str:
    return re.sub(r"\D", "", str(value or ""))


def normalize_match_key(value: str) -> str:
    return re.sub(r"[^а-яa-z0-9]", "", normalize_text(value).lower().replace("ё", "е"))


def unique_list(values: List[str]) -> List[str]:
    result = []
    for value in values:
        if value and value not in result:
            result.append(value)
    return result


def add_unique_values(target: List[str], values: List[str]) -> None:
    for value in values:
        if value and value not in target:
            target.append(value)


def add_unique_funding_rows(target: List[Dict[str, str]], rows: List[Dict[str, str]]) -> None:
    seen = {
        (row.get("Целевая статья"), row.get("КОСГУ"))
        for row in target
    }
    for row in rows:
        key = (row.get("Целевая статья"), row.get("КОСГУ"))
        if row.get("Целевая статья") and row.get("КОСГУ") and key not in seen:
            target.append(row)
            seen.add(key)


def unique_funding_rows(rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    result = []
    add_unique_funding_rows(result, rows)
    return result


def normalize_price_amount(amount: str, kopecks: str = "") -> str:
    amount = normalize_text(amount)

    if re.search(r"[,.]\d{1,2}\b", amount):
        cleaned = re.sub(r"[^\d,.]", "", amount).replace(".", ",")
        rubles, cents = cleaned.split(",", 1)
        rubles = normalize_digits(rubles)
        cents = normalize_digits(cents)[:2].ljust(2, "0")
        return f"{rubles},{cents}" if rubles else ""

    rubles = normalize_digits(amount)
    cents = normalize_digits(kopecks)[:2].zfill(2)
    return f"{rubles},{cents}" if rubles else ""


def extract_money_amounts(text: str) -> List[str]:
    amounts = []
    pattern = re.compile(
        r"(\d[\d\s\u202f]*(?:[,.]\d{1,2})?)"
        r"(?:\s*\([^)]*\))?\s*руб(?:\.|ля|лей)?"
        r"(?:\s*(\d{1,2})\s*коп)?",
        flags=re.IGNORECASE,
    )

    for match in pattern.finditer(normalize_text(text)):
        amount = normalize_price_amount(match.group(1), match.group(2) or "")
        if amount:
            amounts.append(amount)

    return amounts


def price_to_number(value: str) -> float:
    return float(str(value).replace(" ", "").replace(",", "."))


def select_contract_price_amount(amounts: List[str]) -> str:
    if not amounts:
        return ""

    return max(amounts, key=price_to_number)


def extract_contract_price_from_text(text: str) -> str:
    text = normalize_text(text)
    if not text:
        return ""

    patterns = [
        re.compile(
            r"цен[аы]\s+контракта\s+составляет\s+"
            r"(\d[\d\s\u202f]*(?:[,.]\d{1,2})?)"
            r"(?:\s*\([^)]*\))?\s*руб(?:\.|ля|лей)?"
            r"(?:\s*(\d{1,2})\s*коп)?",
            flags=re.IGNORECASE,
        ),
        re.compile(
            r"цен[аы]\s+контракта\s+"
            r"(\d[\d\s\u202f]*(?:[,.]\d{1,2})?)"
            r"(?:\s*\([^)]*\))?\s*руб(?:\.|ля|лей)?"
            r"(?:\s*(\d{1,2})\s*коп)?",
            flags=re.IGNORECASE,
        ),
        re.compile(
            r"цена\s+.*?составляет\s+"
            r"(\d[\d\s\u202f]*(?:[,.]\d{1,2})?)"
            r"(?:\s*\([^)]*\))?\s*руб(?:\.|ля|лей)?"
            r"(?:\s*(\d{1,2})\s*коп)?",
            flags=re.IGNORECASE,
        ),
    ]

    for pattern in patterns:
        match = pattern.search(text)
        if match:
            return normalize_price_amount(match.group(1), match.group(2) or "")

    return ""


def extract_contract_price_from_context(text: str) -> str:
    text = normalize_text(text)
    if not text:
        return ""

    fragments = re.findall(
        r".{0,140}(?:цен[аы]\s+контракта|стоимость\s+контракта|контракт[а-я\s]+цен[аы]).{0,220}",
        text,
        flags=re.IGNORECASE,
    )
    amounts = []
    for fragment in fragments:
        amounts.extend(extract_money_amounts(fragment))

    return select_contract_price_amount(amounts)


def extract_budget_years(point_text: str) -> List[str]:
    point_text = normalize_text(point_text)
    years = []

    budget_fragments = []
    budget_fragments.extend(re.findall(
        r"(?:федеральн\w*\s+)?бюджет\w*.{0,140}?(?:год(?:ы|а)?|г\.)",
        point_text,
        flags=re.IGNORECASE,
    ))
    budget_fragments.extend(re.findall(
        r"финансирован\w*.{0,140}?(?:год(?:ы|а)?|г\.)",
        point_text,
        flags=re.IGNORECASE,
    ))
    budget_fragments.extend(re.findall(
        r"средств\w*.{0,80}?бюджет\w*.{0,100}?(?:год(?:ы|а)?|г\.)",
        point_text,
        flags=re.IGNORECASE,
    ))

    for fragment in budget_fragments:
        years.extend(extract_years_from_text(fragment))

    if not years:
        context_fragments = re.findall(
            r".{0,120}(?:бюджет|финансирован|КОСГУ).{0,120}",
            point_text,
            flags=re.IGNORECASE,
        )
        for fragment in context_fragments:
            years.extend(extract_years_from_text(fragment))

    if not years:
        years = re.findall(r"\b20\d{2}\b(?=\s*(?:финансов\w*\s+)?(?:год|г\.))", point_text, flags=re.IGNORECASE)

    return unique_list(years)


def extract_years_from_text(text: str) -> List[str]:
    years = []
    for start_year, end_year in re.findall(r"\b(20\d{2})\s*[-–—]\s*(20\d{2})\b", text):
        years.extend([start_year, end_year])

    years.extend(re.findall(r"\b20\d{2}\b(?=\s*(?:финансов\w*\s+)?(?:год|г\.))", text, flags=re.IGNORECASE))

    if "бюджет" in text.lower() or "финансирован" in text.lower() or "косгу" in text.lower():
        years.extend(re.findall(r"\b20\d{2}\b", text))

    return unique_list(years)


KOSGU_PATTERN = re.compile(
    r"(?:код\s+)?(?:подстат(?:ья|ьи)\s+)?КОСГУ\s*[:№Nn\-–—]?\s*(\d{3})\b",
    flags=re.IGNORECASE,
)


TECHNICAL_FIELD_ALIASES = {
    "Торговое наименование": [
        "торговое наименование",
        "наименование лекарственного препарата",
    ],
    "Номер регистрационного удостоверения лекарственного препарата": [
        "номер регистрационного удостоверения лекарственного препарата",
        "регистрационное удостоверение",
    ],
    "Количество единиц измерения Товара во вторичной (потребительской) упаковке": [
        "количество единиц измерения товара во вторичной потребительской упаковке",
        "количество единиц измерения во вторичной упаковке",
        "количество единиц во вторичной упаковке",
        "количество во вторичной упаковке",
    ],
    "Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата": [
        "форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата",
        "форма выпуска",
        "лекарственная форма",
    ],
    "Страна происхождения Товара": [
        "страна происхождения товара",
        "страна происхождения",
        "страны происхождения товара",
        "страны происхождения",
    ],
    "Единица измерения": [
        "единица измерения товара",
        "единица измерения",
        "ед изм",
    ],
}


def technical_field_matches(field_key: str, parameter: str) -> bool:
    parameter_key = normalize_match_key(parameter)
    if not parameter_key:
        return False

    for alias in TECHNICAL_FIELD_ALIASES.get(field_key, [field_key]):
        alias_key = normalize_match_key(alias)
        if alias_key and (alias_key in parameter_key or parameter_key in alias_key):
            return True

    return False


def extract_kosgu_codes(text: str) -> List[str]:
    return KOSGU_PATTERN.findall(text or "")


def extract_funding_rows_from_text(text: str) -> List[Dict[str, str]]:
    rows = []
    fragments = re.split(r"\n|;", text or "")

    for fragment in fragments:
        if "косгу" not in fragment.lower():
            continue

        target_match = re.search(
            r"целев[а-яё\s]+стать[а-яё\s]+([\d\s]{5,25})",
            fragment,
            flags=re.IGNORECASE,
        )
        kosgu_codes = extract_kosgu_codes(fragment)
        if not target_match or not kosgu_codes:
            continue

        target_article = normalize_digits(target_match.group(1))
        if target_article:
            rows.append({
                "Целевая статья": target_article,
                "КОСГУ": kosgu_codes[0],
            })

    if rows:
        return rows

    pattern = re.compile(
        r"целев[а-яё\s]+стать[а-яё\s]+([\d\s]{5,25}).{0,160}?"
        r"(?:код\s+)?(?:подстат(?:ья|ьи)\s+)?КОСГУ\s*[:№Nn\-–—]?\s*(\d{3})",
        flags=re.IGNORECASE | re.DOTALL,
    )
    for match in pattern.finditer(text or ""):
        target_article = normalize_digits(match.group(1))
        kosgu_code = match.group(2)
        if target_article and kosgu_code:
            rows.append({
                "Целевая статья": target_article,
                "КОСГУ": kosgu_code,
            })

    return rows


class ContractParser:
    """Парсер для документов контрактов в формате .docx"""

    def __init__(self, docx_path: str):
        self.docx_path = docx_path
        if not os.path.exists(docx_path):
            raise FileNotFoundError(f"Файл не найден: {docx_path}")
        self.doc = Document(docx_path)

    def find_table_by_title(self, title: str) -> List[List[str]]:
        for table in self.doc.tables:
            if table.rows:
                table_data_temp = self.extract_table_data(table)

                if title == "ТЕХНИЧЕСКИЕ ХАРАКТЕРИСТИКИ":
                    for row in table_data_temp[:3]:
                        row_text = ' '.join(row).lower()
                        if (
                            ("№" in row_text and "параметр" in row_text)
                            or ("характеристик" in row_text and "значение" in row_text)
                            or ("наименование показателя" in row_text and "значение" in row_text)
                        ):
                            return table_data_temp

        return []

    def extract_table_data(self, table) -> List[List[str]]:
        table_data = []
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(normalize_text(cell.text))
            if any(cell.strip() for cell in row_data):
                table_data.append(row_data)
        return table_data

    def extract_contract_point(self, point_number: str) -> str:
        lines = collect_document_lines(self.doc)
        start_pattern = re.compile(rf"^\s*{re.escape(point_number)}\.", flags=re.IGNORECASE)

        start_index = -1
        for index, line in enumerate(lines):
            if start_pattern.search(line):
                start_index = index
                break

        if start_index == -1:
            return ""

        next_point_number = point_number.rsplit(".", 1)
        if len(next_point_number) == 2 and next_point_number[1].isdigit():
            next_number = f"{next_point_number[0]}.{int(next_point_number[1]) + 1}"
        else:
            next_number = ""

        next_pattern = re.compile(rf"^\s*{re.escape(next_number)}\.", flags=re.IGNORECASE) if next_number else None
        point_lines = []
        for line in lines[start_index:]:
            if point_lines and next_pattern and next_pattern.search(line):
                break
            point_lines.append(line)

        return "\n".join(point_lines)

    def extract_contract_section(self, section_number: str) -> str:
        lines = collect_document_lines(self.doc)
        start_pattern = re.compile(rf"^\s*{re.escape(section_number)}(?:\.|\s)", flags=re.IGNORECASE)
        next_section = str(int(section_number) + 1) if section_number.isdigit() else ""
        next_pattern = re.compile(rf"^\s*{re.escape(next_section)}(?:\.|\s)", flags=re.IGNORECASE) if next_section else None

        start_index = -1
        for index, line in enumerate(lines):
            if start_pattern.search(line):
                start_index = index
                break

        if start_index == -1:
            return ""

        section_lines = []
        for line in lines[start_index:]:
            if section_lines and next_pattern and next_pattern.search(line):
                break
            section_lines.append(line)

        return "\n".join(section_lines)

    def parse_funding_tables(self) -> Dict[str, Any]:
        result = {
            "years": [],
            "kosgu_codes": [],
            "funding_rows": [],
        }

        for table in self.doc.tables:
            table_data = self.extract_table_data(table)
            if not table_data:
                continue

            table_text = "\n".join(" | ".join(row) for row in table_data)
            table_text_lower = table_text.lower()
            if not any(marker in table_text_lower for marker in ["косгу", "целевая статья", "финансирован", "бюджет"]):
                continue

            add_unique_values(result["years"], extract_budget_years(table_text))
            add_unique_values(result["kosgu_codes"], extract_kosgu_codes(table_text))
            add_unique_funding_rows(result["funding_rows"], extract_funding_rows_from_text(table_text))

            header_index = -1
            target_col = -1
            kosgu_col = -1
            year_col = -1

            for row_index, row in enumerate(table_data):
                for column_index, cell in enumerate(row):
                    cell_lower = cell.lower()
                    if "целевая" in cell_lower and "статья" in cell_lower:
                        target_col = column_index
                    if "косгу" in cell_lower:
                        kosgu_col = column_index
                    if "год" in cell_lower or "бюджет" in cell_lower:
                        year_col = column_index

                if target_col != -1 and kosgu_col != -1:
                    header_index = row_index
                    break

            if header_index == -1:
                continue

            for row in table_data[header_index + 1:]:
                target_article = normalize_digits(row[target_col]) if len(row) > target_col else ""
                kosgu_code = normalize_digits(row[kosgu_col]) if len(row) > kosgu_col else ""
                year_value = row[year_col] if year_col != -1 and len(row) > year_col else ""

                if target_article and len(kosgu_code) >= 3:
                    result["funding_rows"].append({
                        "Целевая статья": target_article,
                        "КОСГУ": kosgu_code[-3:],
                    })
                if len(kosgu_code) == 3:
                    add_unique_values(result["kosgu_codes"], [kosgu_code])
                add_unique_values(result["years"], extract_years_from_text(year_value))

        result["funding_rows"] = unique_funding_rows(result["funding_rows"])
        return result

    def parse_funding_terms(self) -> Dict[str, Any]:
        point_text = self.extract_contract_point("9.1")
        section_text = self.extract_contract_section("9")
        table_terms = self.parse_funding_tables()

        budget_years = extract_budget_years(point_text)
        if not budget_years:
            budget_years = extract_budget_years(section_text)
        if not budget_years:
            budget_years = table_terms["years"]

        kosgu_codes = extract_kosgu_codes(point_text)
        if not kosgu_codes:
            kosgu_codes = extract_kosgu_codes(section_text)
        if not kosgu_codes:
            kosgu_codes = table_terms["kosgu_codes"]

        funding_rows = []
        add_unique_funding_rows(funding_rows, extract_funding_rows_from_text(point_text))
        add_unique_funding_rows(funding_rows, extract_funding_rows_from_text(section_text))
        add_unique_funding_rows(funding_rows, table_terms["funding_rows"])

        result = {}
        if budget_years:
            result["Годы финансирования"] = budget_years
            result["Период поставок по годам"] = "-".join(budget_years)
        if kosgu_codes:
            result["Коды КОСГУ"] = kosgu_codes
            result["Уникальные КОСГУ"] = unique_list(kosgu_codes)
        if funding_rows:
            result["Финансирование КОСГУ"] = funding_rows

        return result

    def extract_contract_price(self) -> str:
        point_text = self.extract_contract_point("2.2")
        price = extract_contract_price_from_text(point_text)
        if price:
            return price
        price = extract_contract_price_from_context(point_text)
        if price:
            return price

        section_text = self.extract_contract_section("2")
        price = extract_contract_price_from_text(section_text)
        if price:
            return price
        price = extract_contract_price_from_context(section_text)
        if price:
            return price

        return self.extract_contract_price_from_tables()

    def extract_contract_price_from_tables(self) -> str:
        amounts = []
        for table in self.doc.tables:
            table_data = self.extract_table_data(table)
            if not table_data:
                continue

            table_text = "\n".join(" | ".join(row) for row in table_data)
            table_text_lower = table_text.lower()
            if not any(marker in table_text_lower for marker in [
                "цена контракта",
                "стоимость контракта",
                "цена государственного контракта",
            ]):
                continue

            price = extract_contract_price_from_text(table_text)
            if price:
                return price
            amounts.extend(extract_money_amounts(table_text))

        return select_contract_price_amount(amounts)

    def parse_contract_terms(self) -> Dict[str, str]:
        price = self.extract_contract_price()
        if price:
            return {"Цена контракта": price}

        return {}

    def find_country_origin_fallback(self, table_data: List[List[str]]) -> str:
        """
        Ищет страну происхождения в мини-таблицах внутри технических характеристик.
        В части контрактов это не строка "Параметр -> Требуемое значение".
        """
        for row_index, row in enumerate(table_data):
            for column_index, cell in enumerate(row):
                cell_lower = cell.lower()
                has_country_header = (
                    'страна происхождения товара' in cell_lower
                    or 'страны происхождения товара' in cell_lower
                )

                if not has_country_header:
                    continue

                current_row_values = [
                    clean_country_origin_value(value)
                    for value in row[column_index + 1:]
                    if is_valid_country_origin(value)
                ]
                current_row_values = unique_values(current_row_values)
                if current_row_values:
                    return '; '.join(current_row_values)

                if row_index + 1 >= len(table_data):
                    continue

                next_row = table_data[row_index + 1]
                if len(next_row) <= column_index:
                    continue

                value = clean_country_origin_value(next_row[column_index])
                if is_valid_country_origin(value):
                    return value

        return ""

    def find_registration_number_fallback(self, table_data: List[List[str]]) -> str:
        for row in table_data:
            values = [
                cell
                for cell in row
                if is_registration_value(cell)
            ]
            if values:
                return format_values(values)

        return ""

    def find_product_table_data(self) -> Dict[str, str]:
        for table in self.doc.tables:
            table_data = self.extract_table_data(table)
            if len(table_data) < 4:
                continue

            header_text = ' '.join(' '.join(row).lower() for row in table_data[:3])
            if (
                'торговое наименование' not in header_text
                or 'количество единиц измерения товара во вторичной' not in header_text
            ):
                continue

            header_rows = table_data[:3]

            def is_numbering_row(row):
                checked_cells = [cell for cell in row[:10] if cell]
                if not checked_cells:
                    return False

                numeric_cells = [
                    cell
                    for cell in checked_cells
                    if cell.rstrip('.').isdigit()
                ]
                return len(numeric_cells) == len(checked_cells)

            data_rows = [
                row
                for row in table_data[3:]
                if row
                and row[0].rstrip('.').isdigit()
                and not is_numbering_row(row)
                and not any('итого' in cell.lower() for cell in row[:2])
            ]

            if not data_rows:
                continue

            def find_column(predicate):
                max_len = max(len(row) for row in header_rows)
                for column_index in range(max_len):
                    texts = []
                    for header_row in header_rows:
                        if len(header_row) > column_index:
                            texts.append(header_row[column_index].lower())
                    if predicate(texts, column_index):
                        return column_index
                return -1

            trade_col = find_column(
                lambda texts, index: any(text == 'торговое наименование' for text in texts)
                and not any('форма выпуска' in text for text in texts)
            )
            form_col = find_column(lambda texts, index: any('торговое наименование, форма выпуска' in text for text in texts))
            unit_col = find_column(lambda texts, index: any(text == 'единица измерения товара' for text in texts))
            quantity_col = find_column(lambda texts, index: any('количество единиц измерения товара во вторичной' in text for text in texts))

            result = {}
            if trade_col != -1:
                result["Торговое наименование"] = format_values([
                    row[trade_col]
                    for row in data_rows
                    if len(row) > trade_col
                ])
            if form_col != -1:
                form_values = []
                for row in data_rows:
                    if len(row) <= form_col:
                        continue
                    value = row[form_col]
                    trade = row[trade_col] if trade_col != -1 and len(row) > trade_col else ""
                    if trade and value.lower().startswith(trade.lower() + ','):
                        value = value[len(trade) + 1:].strip()
                    form_values.append(value)
                result["Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата"] = format_values(form_values)
            if unit_col != -1:
                result["Единица измерения"] = format_values([
                    row[unit_col]
                    for row in data_rows
                    if len(row) > unit_col
                ])
            if quantity_col != -1:
                result["Количество единиц измерения Товара во вторичной (потребительской) упаковке"] = format_values([
                    row[quantity_col]
                    for row in data_rows
                    if len(row) > quantity_col and is_quantity_value(row[quantity_col])
                ])

            return result

        return {}

    def parse_technical_characteristics(self) -> Dict[str, str]:
        table_data = self.find_table_by_title("ТЕХНИЧЕСКИЕ ХАРАКТЕРИСТИКИ")

        if not table_data:
            for table in self.doc.tables:
                if table.rows and len(table.rows) > 1:
                    table_data_temp = self.extract_table_data(table)
                    for row in table_data_temp[:3]:
                        row_text = ' '.join(row)
                        if "№" in row_text and "Параметр" in row_text:
                            table_data = table_data_temp
                            break
                    if table_data:
                        break

        needed_fields = {
            "Торговое наименование": "",
            "Номер регистрационного удостоверения лекарственного препарата": "",
            "Количество единиц измерения Товара во вторичной (потребительской) упаковке": "",
            "Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата": "",
            "Страна происхождения Товара": "",
            "Единица измерения": "",
        }

        if table_data:
            column_param = -1
            column_value = -1
            start_index = 1

            for i, row in enumerate(table_data):
                if len(row) >= 2:
                    param_columns = []
                    value_columns = []
                    for j, cell in enumerate(row):
                        cell_lower = cell.lower()
                        if 'параметр' in cell_lower:
                            param_columns.append(j)
                        elif 'требуемое значение' in cell_lower or 'значение' in cell_lower:
                            value_columns.append(j)

                    if param_columns and value_columns:
                        column_param = param_columns[0]
                        column_value = next(
                            (column for column in value_columns if column > column_param),
                            value_columns[0]
                        )
                        start_index = i + 1
                        break

            if column_param == -1 or column_value == -1:
                column_param = 1
                column_value = 2
                start_index = 1

            for i in range(start_index, len(table_data)):
                row = table_data[i]
                if not row or not any(cell.strip() for cell in row):
                    continue

                if len(row) > column_param and len(row) > column_value:
                    parameter = row[column_param].strip()
                    value = row[column_value].strip()

                    if not parameter or not value:
                        continue

                    if parameter[:2].replace('.', '').isdigit():
                        parameter = parameter.split('.', 1)[1].strip() if '.' in parameter else parameter

                    for needed_param in needed_fields.keys():
                        if technical_field_matches(needed_param, parameter):
                            if not needed_fields[needed_param]:
                                needed_fields[needed_param] = value
                            break

            if not needed_fields["Страна происхождения Товара"]:
                needed_fields["Страна происхождения Товара"] = self.find_country_origin_fallback(table_data)

            if not is_valid_country_origin(needed_fields["Страна происхождения Товара"]):
                needed_fields["Страна происхождения Товара"] = self.find_country_origin_fallback(table_data)

            if not is_registration_value(needed_fields["Номер регистрационного удостоверения лекарственного препарата"]):
                needed_fields["Номер регистрационного удостоверения лекарственного препарата"] = self.find_registration_number_fallback(table_data)

            product_table_data = self.find_product_table_data()
            if product_table_data:
                if not is_quantity_value(needed_fields["Количество единиц измерения Товара во вторичной (потребительской) упаковке"]):
                    needed_fields["Количество единиц измерения Товара во вторичной (потребительской) упаковке"] = product_table_data.get(
                        "Количество единиц измерения Товара во вторичной (потребительской) упаковке",
                        needed_fields["Количество единиц измерения Товара во вторичной (потребительской) упаковке"]
                    )

                form_value = needed_fields["Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата"]
                if form_value and '(' not in form_value and product_table_data.get("Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата"):
                    needed_fields["Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата"] = product_table_data[
                        "Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата"
                    ]

                if product_table_data.get("Единица измерения"):
                    needed_fields["Единица измерения"] = product_table_data["Единица измерения"]

        result = {}
        for key, value in needed_fields.items():
            if value:
                if key == "Страна происхождения Товара":
                    value = clean_country_origin_value(value)
                elif key == "Единица измерения":
                    value = normalize_unit_value(value)
                elif key == "Форма выпуска в соответствии с регистрационным удостоверением лекарственного препарата":
                    value = normalize_form_value(value)
                result[key] = value

        return result

    def extract_all_data(self) -> Dict[str, Any]:
        return {
            "technical_characteristics": self.parse_technical_characteristics(),
            "delivery_terms": parse_delivery_terms(self.doc),
            "funding_terms": self.parse_funding_terms(),
            "contract_terms": self.parse_contract_terms(),
        }


def parse_contract(contract_file: str) -> Dict[str, Any]:
    return ContractParser(contract_file).extract_all_data()
