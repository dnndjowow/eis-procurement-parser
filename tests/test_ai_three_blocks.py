import contextlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

BASE = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


engine = load('three_engine', BASE / 'test.py')
three = load('three_blocks', BASE / 'ai_three_blocks.py')


def document():
    doc = Document()
    doc.add_paragraph('На поставку лекарственного препарата')
    doc.add_paragraph('2.2. Цена Контракта составляет 10 руб. 00 коп.')
    doc.add_paragraph('3.1. Иные условия')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Количество вторичных упаковок по потребности'
    table.cell(0, 1).text = 'Цена за единицу (шт.), итого, руб.'
    table.cell(1, 0).text = '10'
    table.cell(1, 1).text = '1'
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Торговое наименование'
    table.cell(0, 1).text = 'Препарат'
    table.cell(1, 0).text = 'Номер регистрационного удостоверения'
    table.cell(1, 1).text = 'РУ-123'
    table = doc.add_table(rows=2, cols=3)
    for col, value in enumerate(('Этап поставки', 'Срок поставки', 'Количество')):
        table.cell(0, col).text = value
    for col, value in enumerate(('I этап', 'до 01.02.2027', '10')):
        table.cell(1, col).text = value
    doc.add_paragraph('9.4. Документы не позднее 01.03.2027')
    return doc


class ThreeBlockTests(unittest.TestCase):
    def test_total_contract_price_source_prefers_exact_clause_22_amount(self):
        sources = engine.DocumentSourceMap({
            'DOC/P21': '2.2. Цена Контракта составляет 55 865 940 руб. 00 коп.',
            'DOC/P22': '2.2.1. Стоимость Товара по 1 этапу 15 564 640 руб. 00 коп.',
        })
        parsed = {'contract_price': '55865940.00', 'sources': ['DOC/P22']}
        corrections = []

        changed = three._repair_contract_price_source(engine, parsed, sources, corrections)

        self.assertTrue(changed)
        self.assertEqual(parsed['sources'], ['DOC/P21'])
        self.assertIn('2.2.1', corrections[0]['reason'])

    def test_missing_product_unit_is_filled_from_exact_labeled_word_row(self):
        sources = engine.DocumentSourceMap({
            'DOC/T1/R1C1': '8.1',
            'DOC/T1/R1C2:R1C3': 'Единица измерения',
            'DOC/T1/R1C4:R1C5': 'МЕ',
            'DOC/T1/R2C1': '8.',
            'DOC/T1/R2C2:R2C3': 'Единица измерения в соответствии с ЕСКЛП',
            'DOC/T1/R2C4:R2C5': 'МЕ',
            'DOC/T1/R3C1': '2.',
            'DOC/T1/R3C2': 'Торговое наименование',
            'DOC/T1/R3C3': 'Препарат А',
            'DOC/T1/R4C1': '4.',
            'DOC/T1/R4C2': 'Номер регистрационного удостоверения',
            'DOC/T1/R4C3': 'РУ-123',
        })
        parsed = {'technical_variants': [{
            'trade_name': 'Препарат А', 'product_unit': None,
            'sources': {'trade_name': ['DOC/T1/R3C3']},
        }], 'warnings': [
            'Поле «единица товара» не выделено; значение не извлечено.',
        ]}
        corrections, notices = [], ['MISSING $.technical_variants[0].product_unit: значение отсутствует']

        three._explicit_single_variant_unit(engine, parsed, sources, corrections, notices)

        variant = parsed['technical_variants'][0]
        self.assertEqual(variant['product_unit'], 'МЕ')
        self.assertEqual(variant['sources']['product_unit'], ['DOC/T1/R1C4:R1C5'])
        self.assertEqual(notices, [])
        self.assertEqual(parsed['warnings'], [])
        self.assertEqual(corrections[0]['source'], 'DOC/T1/R1C4:R1C5')

    def test_product_unit_autofill_does_not_guess_for_multiple_variants(self):
        sources = engine.DocumentSourceMap({
            'DOC/T1/R1C1': '8.1', 'DOC/T1/R1C2': 'Единица измерения',
            'DOC/T1/R1C3': 'МЕ', 'DOC/T1/R2C1': '2.',
            'DOC/T1/R2C2': 'Торговое наименование', 'DOC/T1/R2C3': 'Препарат А',
            'DOC/T1/R3C1': '4.', 'DOC/T1/R3C2': 'Торговое наименование',
            'DOC/T1/R3C3': 'Препарат Б',
        })
        parsed = {'technical_variants': [
            {'product_unit': None, 'sources': {'trade_name': ['DOC/T1/R2C3']}},
            {'product_unit': None, 'sources': {'trade_name': ['DOC/T1/R3C3']}},
        ]}
        corrections, notices = [], []

        three._explicit_single_variant_unit(engine, parsed, sources, corrections, notices)

        self.assertIsNone(parsed['technical_variants'][0]['product_unit'])
        self.assertEqual(corrections, [])

    def test_equivalent_volume_and_decimal_separator_warnings_are_removed(self):
        parsed = {'technical_variants': [{
            'release_form': 'концентрат 4.5 мл х 1 (пачка картонная)',
            'package_contents': [{'quantity': '4.5', 'unit': 'мл'}],
        }], 'warnings': [
            'Значение количества во вторичной упаковке "4,5" не совпадает с "х 1"; возможно, это объём в мл.',
            'Дозировка "4.5 мг/мл" и "4,5 мг/мл" различается из-за разделителя.',
            'Количество 4 мл не совпадает с указанным в другой позиции.',
        ]}
        corrections = []

        three._clean_technical_warnings(engine, parsed, corrections)

        self.assertEqual(parsed['warnings'], ['Количество 4 мл не совпадает с указанным в другой позиции.'])
        self.assertEqual(len(corrections), 2)

    def test_ambiguous_thousands_separator_warning_is_not_removed(self):
        parsed = {'technical_variants': [], 'warnings': [
            'Значения "4.500 мг/мл" и "4,500 мг/мл" различаются из-за разделителя.',
        ]}
        corrections = []

        three._clean_technical_warnings(engine, parsed, corrections)

        self.assertEqual(len(parsed['warnings']), 1)
        self.assertEqual(corrections, [])

    def test_reserve_attempt_cannot_replace_better_primary_result(self):
        primary = {'status': 'REVIEW', 'errors': [], 'notices': ['REVIEW: inspect'],
                   'history': [], 'total_seconds': 2}
        fallback = {'status': 'INVALID', 'errors': ['invalid schema'], 'notices': [],
                    'history': [], 'total_seconds': 3}

        selected = three._select_attempt(engine, primary, fallback)

        self.assertIs(selected, primary)
        self.assertEqual(selected['history'], [fallback])
        self.assertEqual(selected['total_seconds'], 5)

    def test_better_reserve_attempt_is_selected(self):
        primary = {'status': 'INVALID', 'errors': ['invalid schema'], 'notices': [],
                   'history': [], 'total_seconds': 2}
        fallback = {'status': 'OK', 'errors': [], 'notices': [], 'history': [],
                    'total_seconds': 3}

        selected = three._select_attempt(engine, primary, fallback)

        self.assertIs(selected, fallback)
        self.assertEqual(selected['history'], [primary])
        self.assertEqual(selected['total_seconds'], 5)

    def test_numeric_normalization_is_lossless(self):
        cases = [('402.00', 'integer', '402'), ('4,5', 'number', '4.5'),
                 ('1\u00a0234,50', 'number', '1234.50'),
                 ('123.12345678901234567890', 'number', '123.12345678901234567890'),
                 (402, 'integer', '402')]
        for value, kind, expected in cases:
            self.assertEqual(engine.normalize_numeric_value(value, kind), expected)
        for value in ('402.5', '4,500', '12 34', '402 упаковки', '=400+2', '-2', '1e3'):
            self.assertEqual(engine.normalize_numeric_value(value, 'integer'), value)
        self.assertIs(engine.normalize_numeric_value(True, 'integer'), True)
        self.assertEqual(engine.normalize_numeric_value(4.5, 'number'), 4.5)

    def test_numeric_normalization_is_logged(self):
        data = engine.build_record_template(three.schemas(engine)[0])
        data['products'][0]['secondary_packages_whole'] = '402.00'
        changes = []
        parsed, errors = three.validate(engine, json.dumps(data), 0, {}, {}, changes, [])
        self.assertEqual(parsed['products'][0]['secondary_packages_whole'], '402')
        self.assertTrue(any(c['before'] == '402.00' and c['after'] == '402' for c in changes))
        self.assertTrue(errors)  # Missing source is not excused by normalization.

    def test_confirmed_deadline_is_bound_to_document_and_clause(self):
        import hashlib
        text = '9.4. Документы позднее «15» февраля 2027 г.'
        data = {'document_deadline': '2027-02-15', 'stage': 'I этап',
                'sources': {'document_deadline': ['DOC/P1']}}
        source = engine.DocumentSourceMap({'DOC/P1': text}, 'trusted-document')
        key = ('trusted-document', hashlib.sha256(text.encode()).hexdigest())
        with patch.object(engine, 'CONFIRMED_DOCUMENT_DEADLINES', {key: '2027-02-15'}):
            errors, changes = [], []
            engine.check_extraction_evidence(data, 0, source, errors, [], changes)
            self.assertEqual(errors, [])
            self.assertEqual(len(changes), 1)
            for changed in (engine.DocumentSourceMap(source, 'other-document'),
                            engine.DocumentSourceMap({'DOC/P1': text + ' изменено'}, 'trusted-document')):
                errors = []
                engine.check_extraction_evidence(data, 0, changed, errors, [])
                self.assertTrue(any('не подтверждена' in e for e in errors))

    def test_complete_release_form_required(self):
        sources = {'DOC/T1/R1C1': 'Торговое наименование',
                   'DOC/T1/R2C1': 'Форма выпуска в соответствии с регистрационным удостоверением',
                   'DOC/T1/R2C2': '[раствор, 500 МЕ (флакон) х 1] х 1 (пачка картонная)'}
        full = sources['DOC/T1/R2C2']
        for value, valid in ((full, True), (full[1:], False), (full[:-1], False),
                             (full.replace('500', '50'), False),
                             (full.replace(' ', '\n', 1), True)):
            errors = []
            engine.check_extraction_evidence({'release_form': value,
                'sources': {'release_form': ['DOC/T1/R2C2']}}, 0, sources, errors, [])
            self.assertEqual(not errors, valid, value)

    def test_numbered_release_variants_keep_internal_punctuation(self):
        text = '1. [капсулы] (флакон) 120 х 1;\n2. таблетки (упаковка) 10 х 3.'
        candidates = engine.release_form_candidates(text)
        self.assertIn('[капсулы] (флакон) 120 х 1', candidates)
        self.assertIn('таблетки (упаковка) 10 х 3', candidates)
        self.assertNotIn('капсулы] (флакон) 120 х 1', candidates)
        self.assertNotIn('таблетки', candidates)

    def test_full_run_includes_all_three_blocks(self):
        calls = []
        def extracted(engine_arg, model, part, prompt, block):
            calls.append(part)
            return {'model': model, 'part': part + 1, 'status': 'OK', 'seconds': 0,
                    'total_seconds': 0, 'history': [], 'parsed': {}, 'errors': [], 'notices': []}
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docs = root / 'docs'
            docs.mkdir()
            document().save(docs / 'sample.docx')
            with patch.object(three, 'extract', side_effect=extracted), contextlib.redirect_stdout(io.StringIO()):
                out = three.run(engine, docs, root / 'reports')
            self.assertEqual(calls, [0, 1, 2])
            bundle = json.loads((out / '001' / 'bundle.json').read_text())
            self.assertEqual(set(bundle['blocks']), set(three.NAMES))
            self.assertEqual(len(list((out / '001').glob('attempt_*.json'))), 3)

    def test_sections_are_disjoint_and_ids_preserved(self):
        blocks = three.document_blocks(engine, document())
        self.assertEqual(len(blocks), 3)
        self.assertIn('DOC/T1/R2C1', blocks[0][1])
        self.assertIn('DOC/T2/R2C2', blocks[1][1])
        self.assertIn('DOC/T3/R2C3', blocks[2][1])
        for i in range(3):
            for j in range(i):
                self.assertFalse(set(blocks[i][1]) & set(blocks[j][1]))
        self.assertNotIn('9.4.', blocks[0][0])
        self.assertNotIn('2.2.', blocks[1][0])
        self.assertNotIn('Цена за единицу', blocks[1][0])

    def test_schemas_do_not_mix_domains(self):
        spec, tech, delivery = three.schemas(engine)
        self.assertNotIn('technical_variants', spec['products'][0])
        self.assertNotIn('products', tech)
        self.assertNotIn('stages', spec)
        self.assertIn('stages', delivery)
        self.assertEqual(len(three.prompts(engine)), 3)

    def test_ambiguous_table_rejected(self):
        doc = document()
        doc.tables[1].cell(1, 1).text = 'Количество вторичных упаковок по потребности'
        with self.assertRaisesRegex(ValueError, 'Смешанная таблица'):
            three.document_blocks(engine, doc)

    def test_prepare_does_not_call_api(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docs = root / 'docs'
            docs.mkdir()
            document().save(docs / 'sample.docx')
            with patch.object(engine, 'ask_model_with_deadline') as api, contextlib.redirect_stdout(io.StringIO()):
                out = three.run(engine, docs, root / 'reports', prepare_only=True)
            api.assert_not_called()
            self.assertEqual(len(list((out / '001').glob('input_*.txt'))), 3)

    def test_products_mode_sends_two_blocks_without_linking(self):
        calls = []
        def answer(model, prompt):
            part = len(calls)
            calls.append(prompt)
            return json.dumps(engine.build_record_template(three.schemas(engine)[part]))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            docs = root / 'docs'
            docs.mkdir()
            document().save(docs / 'sample.docx')
            with patch.object(engine, 'ask_model_with_deadline', side_effect=answer), \
                 patch.object(engine, 'check_extraction_evidence'), contextlib.redirect_stdout(io.StringIO()):
                out = three.run(engine, docs, root / 'reports', products_only=True)
            self.assertEqual(len(calls), 2)
            bundle = json.loads((out / '001' / 'bundle.json').read_text())
            self.assertEqual(bundle['product_links'], [])
            self.assertEqual(bundle['link_status'], 'NOT_LINKED')
            self.assertTrue(bundle['requires_review'])


if __name__ == '__main__':
    unittest.main()
