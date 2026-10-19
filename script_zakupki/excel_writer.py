from dataclasses import dataclass, field
from pathlib import Path
from copy import copy
import re

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .common import extract_notice_number, normalize_header, values_equal
from .settings import DATE_FORMAT, TARGET_YEAR


FORMULA_COLUMNS_2026 = {
    "AA": "=((W{row}-AC{row})/W{row})*100",
    "AB": "=W{row}-AC{row}",
    "AD": "=W{row}-AE{row}",
    "AF": "=AE{row}",
    "AG": "=AF{row}",
    "AL": "=AF{row}-(AV{row}+AX{row})",
    "AM": "=AG{row}/AP{row}",
    "AN": "=AF{row}/AP{row}",
    "AO": "=AN{row}*BU{row}",
    "AP": "=SUM(AQ{row}:AT{row})",
    "AV": "=AU{row}*AN{row}",
    "AX": "=AW{row}*AN{row}",
    "AY": "=AP{row}/BU{row}",
    "AZ": "=_xlfn.CEILING.MATH(AY{row})",
    "BX": "=(W{row}*10)/100",
    "CL": "=CM{row}-CK{row}",
    "CM": "=AF{row}",
    "CQ": "=AQ{row}*AN{row}",
    "CS": "=AN{row}*AR{row}",
    "CU": "=AN{row}*AS{row}",
    "CW": "=AN{row}*AT{row}",
}

FORMULA_COLUMNS_2027 = {
    "AA": "=((W{row}-AC{row})/W{row})*100",
    "AB": "=W{row}-AC{row}",
    "AD": "=W{row}-AE{row}",
    "AF": "=AE{row}",
    "AG": "=AF{row}",
    "AL": "=AF{row}-(AW{row}+AY{row})",
    "AM": "=AG{row}/AP{row}",
    "AN": "=AF{row}/AP{row}",
    "AO": "=AN{row}*BY{row}",
    "AP": "=SUM(AQ{row}:AT{row})",
    "AW": "=AV{row}*AN{row}",
    "AY": "=AX{row}*AN{row}",
    "AZ": "=AP{row}/BY{row}",
    "BA": "=_xlfn.CEILING.MATH(AZ{row})",
    "CB": "=(W{row}*10)/100",
    "CP": "=CQ{row}-CO{row}",
    "CQ": "=AF{row}",
    "CU": "=AQ{row}*AN{row}",
    "CW": "=AN{row}*AR{row}",
    "CY": "=AN{row}*AS{row}",
    "DA": "=AN{row}*AT{row}",
    "DC": "=AN{row}*AU{row}",
}

FORMULA_COLUMNS_BY_YEAR = {
    2026: FORMULA_COLUMNS_2026,
    2027: FORMULA_COLUMNS_2027,
}

DATE_COLUMNS_BY_YEAR = {
    2026: {
        "E", "G", "H", "AK",
        "BA", "BB", "BC", "BD",
        "BE", "BG", "BI", "BK",
    },
    2027: {
        "E", "G", "H", "AK",
        "BB", "BC", "BD", "BE", "BF",
        "BG", "BI", "BK", "BM", "BO",
    },
}
YEAR_SHEETS = [f"{year} год" for year in range(2026, 2030)]
DATA_ROW_STYLE_TEMPLATE_ROW = 220
GENERAL_FORMAT_COLUMNS_BY_YEAR = {
    2027: {"CH"},
}
DATE_FORMAT_COLUMNS_BY_YEAR = {
    2027: {"CX"},
}
NO_FILL = PatternFill(fill_type=None)


HEADERS = {
    "A": "Бюджет",
    "B": "Период поставок по годам",
    "C": "ИКЗ",
    "D": "Основание для заключения\nНомер извещения/ распоряжение/ ст. № 44-ФЗ",
    "E": "Дата извещения",
    "G": "Дата окончания срока подачи заявок",
    "H": "Дата подведения итогов определения поставщика",
    "P": "Нормативно-правовой акт, в рамках которого осуществляется централизованная закупка лекарственных препаратов",
    "R": "Ссылка на ЕИС",
    "V": "МНН закупаемого лекарственного препарата",
    "W": "НМЦК (руб.)",
    "X": "НМЦК (руб.) первый год",
    "Y": "НМЦК (руб.)   Второй год",
    "Z": "НМЦК (руб.)   Третий год",
    "AE": "цена по контракту на 2026 год",
    "AH": "КБК 2026 (основной)",
    "AI": "КБК 2026 (дети)",
    "AJ": "ОКПД 2",
    "AK": "срок действия ГК",
    "AQ": "I эт",
    "AR": "II эт",
    "AS": "III эт",
    "AT": "IV эт",
    "AU": "Итого кол-во 2026 дети",
    "AW": "Итого кол-во 2026 взрослые",
    "BA": "Срок поставки по ГК 2026\nI этап",
    "BB": "Срок поставки по ГК 2026\nII этап",
    "BC": "Срок поставки по ГК 2026\nIII этап",
    "BD": "Срок поставки по ГК 2026\nIV этап",
    "BE": "Срок предоставления документов (план)\nI этап",
    "BG": "Срок предоставления документов (план)\nII этап",
    "BI": "Срок предоставления документов (план)\nIII этап",
    "BK": "Срок предоставления документов (план)\nIV этап",
    "BW": "Обеспечение исп. ГК\n% ",
    "CK": "цена оплаченная по контракту",
}


@dataclass
class ExcelChange:
    cell: str
    old_value: object
    new_value: object


@dataclass
class ExcelWriteResult:
    row: int
    changed: bool
    changes: list[ExcelChange] = field(default_factory=list)


def load_or_create_workbook(path):
    path = Path(path)
    if path.exists():
        return load_workbook(path)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = YEAR_SHEETS[0]
    for column, header in HEADERS.items():
        sheet[f"{column}1"] = header
    for column in FORMULA_COLUMNS_2026:
        if not sheet[f"{column}1"].value:
            sheet[f"{column}1"] = column
    return workbook


def get_excel_lock_path(workbook_path):
    workbook_path = Path(workbook_path)
    return workbook_path.with_name(f"~${workbook_path.name}")


def ensure_excel_is_closed(workbook_path):
    lock_path = get_excel_lock_path(workbook_path)
    if lock_path.exists():
        raise RuntimeError(f"закрой Excel-файл перед запуском: найден {lock_path.name}")


def get_year_sheet(workbook, year):
    sheet_name = f"{year} год"
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]

    return None


def get_year_sheets(workbook):
    return {
        int(sheet_name.split()[0]): workbook[sheet_name]
        for sheet_name in YEAR_SHEETS
        if sheet_name in workbook.sheetnames
    }


def clear_publication_data_rows(sheet):
    if sheet.max_row <= 1:
        return 0

    cleared_rows = 0
    clear_fill = get_sheet_year(sheet) == 2027
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
        row_has_value = any(cell.value not in (None, "") for cell in row)
        if row_has_value:
            cleared_rows += 1

        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.value = None
            if clear_fill:
                cell.fill = copy(NO_FILL)

    return cleared_rows


def find_or_create_row(sheet, notice_number):
    notice_col = "D"
    for row in range(2, sheet.max_row + 1):
        current_notice = extract_notice_number(sheet[f"{notice_col}{row}"].value)
        if current_notice == notice_number:
            return row

    return find_append_row(sheet)


def find_append_row(sheet):
    if get_sheet_year(sheet) == 2027:
        key_columns = ["B", "C", "D", "E", "G", "H", "P", "V", "W", "AQ", "BB", "BG"]
    else:
        key_columns = ["B", "C", "D", "E", "G", "H", "P", "V", "W", "AQ", "BA", "BE"]
    last_data_row = 1

    for row in range(2, sheet.max_row + 1):
        if any(sheet[f"{column}{row}"].value not in (None, "") for column in key_columns):
            last_data_row = row

    return last_data_row + 1


def write_purchase_row(sheet, data, apply_changes):
    row = find_append_row(sheet) if data.get("_append_row") else find_or_create_row(sheet, data["D"])
    changes = []

    if apply_changes:
        apply_data_row_style(sheet, row)

    values_to_write = {
        column: value
        for column, value in data.items()
        if not column.startswith("_") and value not in (None, "")
    }

    if is_multiyear_data(data):
        values_to_write["AE"] = None

    for column in data.get("_clear_columns", []):
        values_to_write.setdefault(column, None)

    for column, formula in build_formula_columns(sheet, data).items():
        values_to_write[column] = formula.format(row=row)

    for column, value in sorted(values_to_write.items(), key=lambda item: column_index(item[0])):
        value = render_row_value(value, row)
        cell = sheet[f"{column}{row}"]
        old_value = cell.value

        if values_equal(old_value, value):
            continue

        changes.append(ExcelChange(cell.coordinate, old_value, value))
        if apply_changes:
            cell.value = value
            if is_date_column(sheet, column):
                cell.number_format = DATE_FORMAT

    return ExcelWriteResult(row=row, changed=bool(changes), changes=changes)


def apply_data_row_style(sheet, row, template_row=DATA_ROW_STYLE_TEMPLATE_ROW):
    if template_row > sheet.max_row:
        template_row = 2

    if row == template_row:
        return

    template_height = sheet.row_dimensions[template_row].height
    if template_height is not None:
        sheet.row_dimensions[row].height = template_height

    copy_fill = get_sheet_year(sheet) != 2027

    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(template_row, column)
        target = sheet.cell(row, column)
        if isinstance(target, MergedCell):
            continue

        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill) if copy_fill else copy(NO_FILL)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)

        if is_general_format_column(sheet, get_column_letter(column)):
            target.number_format = "General"
        if is_forced_date_format_column(sheet, get_column_letter(column)):
            target.number_format = DATE_FORMAT


def build_formula_columns(sheet, data):
    formulas = dict(get_formula_columns(sheet))

    if get_sheet_year(sheet) == 2027 and data.get("AU") not in (None, ""):
        formulas["AP"] = "=SUM(AQ{row}:AU{row})"

    if is_multiyear_data(data):
        nmc_column = resolve_current_year_nmc_column(sheet, data)
        formulas["AA"] = f"=(({nmc_column}{{row}}-AC{{row}})/{nmc_column}{{row}})*100"
        formulas["AB"] = f"={nmc_column}{{row}}-AC{{row}}"
        formulas["AD"] = f"={nmc_column}{{row}}-AE{{row}}"

    for column in data.get("_manual_formula_columns", []):
        formulas.pop(column, None)

    return formulas


def get_formula_columns(sheet):
    return FORMULA_COLUMNS_BY_YEAR.get(get_sheet_year(sheet), FORMULA_COLUMNS_2026)


def is_date_column(sheet, column):
    return column in DATE_COLUMNS_BY_YEAR.get(get_sheet_year(sheet), DATE_COLUMNS_BY_YEAR[2026])


def is_general_format_column(sheet, column):
    return column in GENERAL_FORMAT_COLUMNS_BY_YEAR.get(get_sheet_year(sheet), set())


def is_forced_date_format_column(sheet, column):
    return column in DATE_FORMAT_COLUMNS_BY_YEAR.get(get_sheet_year(sheet), set())


def is_multiyear_data(data):
    return len(get_financing_years(data)) > 1


def resolve_current_year_nmc_column(sheet, data):
    target_year = get_target_year_from_contract_price_header(sheet)
    years = get_financing_years(data)
    year_columns = ["X", "Y", "Z"]

    if target_year in years:
        year_index = years.index(target_year)
        if year_index < len(year_columns):
            return year_columns[year_index]

    return "W"


def get_target_year_from_contract_price_header(sheet):
    header = str(sheet["AE1"].value or "")
    match = re.search(r"на\s+(20\d{2})\s+год", header, flags=re.IGNORECASE)
    if match:
        return int(match.group(1))

    return TARGET_YEAR


def get_sheet_year(sheet):
    match = re.search(r"\b(20\d{2})\b", str(sheet.title))
    if match:
        return int(match.group(1))

    return TARGET_YEAR


def get_financing_years(data):
    years = data.get("_financing_years")
    if years:
        return [int(year) for year in years]

    period = str(data.get("B") or "")
    return [int(year) for year in re.findall(r"20\d{2}", period)]


def render_row_value(value, row):
    if isinstance(value, str) and "{row}" in value:
        return value.format(row=row)
    return value


def column_index(column):
    result = 0
    for char in column:
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def find_header_columns(sheet):
    result = {}
    for cell in sheet[1]:
        if cell.value not in (None, ""):
            result[normalize_header(cell.value)] = get_column_letter(cell.column)
    return result
