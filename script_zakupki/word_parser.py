import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document

from .common import (
    MONTHS,
    extract_ikz,
    extract_notice_number,
    normalize_text,
    parse_all_dates,
    parse_number,
)


@dataclass
class DeliveryStage:
    stage: str
    delivery_date: object = None
    quantity: object = None
    children_quantity: object = None
    adult_quantity: object = None

    @property
    def total_quantity(self):
        if self.quantity is not None:
            return self.quantity

        total = 0
        has_value = False
        for value in (self.children_quantity, self.adult_quantity):
            if value is not None:
                total += value
                has_value = True
        return total if has_value else None


@dataclass
class WordPurchaseData:
    path: Path
    notice_number: str | None = None
    ikz: str | None = None
    drug_name: str | None = None
    contract_unit_price: object = None
    package_quantity: object = None
    budget_years: list[int] = field(default_factory=list)
    stages: list[DeliveryStage] = field(default_factory=list)
    delivery_stage_count: int = 0
    delivery_stage_years: list[int] = field(default_factory=list)
    document_deadlines: dict[str, object] = field(default_factory=dict)


def iter_docx_files(files_dir):
    files_dir = Path(files_dir)
    if not files_dir.exists():
        return []

    return sorted(
        path
        for path in files_dir.glob("*.docx")
        if not path.name.startswith("~$")
    )


def parse_purchase_word(path):
    path = Path(path)
    doc = Document(path)
    text = collect_document_text(doc)
    delivery_stage_summary = summarize_delivery_stages(doc)

    return WordPurchaseData(
        path=path,
        notice_number=extract_notice_number(text),
        ikz=extract_ikz(text),
        drug_name=parse_drug_name(text),
        contract_unit_price=parse_contract_unit_price(doc),
        package_quantity=parse_package_quantity(doc),
        budget_years=parse_budget_years(doc),
        stages=parse_delivery_stages(doc),
        delivery_stage_count=delivery_stage_summary["count"],
        delivery_stage_years=delivery_stage_summary["years"],
        document_deadlines=parse_document_deadlines(doc),
    )


def collect_document_text(doc):
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def collect_document_lines(doc):
    lines = []
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(normalize_text(cell.text) for cell in row.cells)
            if row_text.strip(" |"):
                lines.append(row_text)
    return lines


def parse_budget_years(doc):
    clause = extract_clause_text(doc, "9.1.")
    if not clause:
        return []

    years = sorted({int(year) for year in re.findall(r"\b20\d{2}\b", clause)})
    return years


def parse_drug_name(text):
    text = normalize_text(text)
    match = re.search(
        r"Поставка иммунобиологического лекарственного препарата\s+для иммунопрофилактики\s+(.+?)\s+в рамках",
        text,
        flags=re.IGNORECASE,
    )
    if match:
        return normalize_drug_name(match.group(1))

    match = re.search(
        r"Поставка иммунобиологического лекарственного препарата\s+(.+?)\s+в рамках",
        text,
        flags=re.IGNORECASE,
    )
    if match:
        return normalize_drug_name(match.group(1))

    match = re.search(
        r"на поставку лекарственного препарата\s+(.+?)\s+в рамках",
        text,
        flags=re.IGNORECASE,
    )
    if match:
        return normalize_drug_name(match.group(1))

    match = re.search(r"Поставка лекарственного препарата\s+(.+?)\s+в рамках", text, flags=re.IGNORECASE)
    if match:
        return normalize_drug_name(match.group(1))

    return None


def normalize_drug_name(value):
    text = normalize_text(value).strip(" .")
    if not text:
        return None
    return text[:1].upper() + text[1:]


def parse_stage(value):
    text = normalize_text(value).lower()
    exact_value = text.strip(" .)")
    standalone_stages = {
        "i": "I",
        "1": "I",
        "ii": "II",
        "2": "II",
        "iii": "III",
        "3": "III",
        "iv": "IV",
        "4": "IV",
        "v": "V",
        "5": "V",
        "vi": "VI",
        "6": "VI",
    }
    if exact_value in standalone_stages:
        return standalone_stages[exact_value]

    match = re.search(
        r"\b(i{1,3}|iv|v|vi|1|2|3|4|5|6|перв\w*|втор\w*|трет\w*|четвер\w*|пят\w*|шест\w*)\s+этап",
        text,
    )
    if not match:
        return None

    token = match.group(1)
    mapping = {
        "i": "I",
        "1": "I",
        "ii": "II",
        "2": "II",
        "iii": "III",
        "3": "III",
        "iv": "IV",
        "4": "IV",
        "v": "V",
        "5": "V",
        "vi": "VI",
        "6": "VI",
    }
    if token in mapping:
        return mapping[token]

    for prefix, stage in (
        ("перв", "I"),
        ("втор", "II"),
        ("трет", "III"),
        ("четвер", "IV"),
        ("пят", "V"),
        ("шест", "VI"),
    ):
        if token.startswith(prefix):
            return stage

    return None


def parse_delivery_stages(doc):
    stages = []

    for table in doc.tables:
        rows = [
            [normalize_text(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        header_index = find_delivery_header_index(rows)
        if header_index is None:
            continue

        header_rows = rows[header_index:header_index + 3]
        stage_col, date_col, quantity_cols, children_col, adult_col = detect_delivery_columns(header_rows)
        if stage_col is None or date_col is None:
            continue

        for row in rows[header_index + 1:]:
            if len(row) <= max(stage_col, date_col):
                continue

            stage = parse_stage(row[stage_col])
            if not stage:
                continue

            dates = parse_all_dates(row[date_col])
            delivery_date = dates[-1] if dates else None
            quantity = parse_generic_stage_quantity(row, quantity_cols)
            children_quantity = parse_stage_quantity(row[children_col]) if children_col is not None and len(row) > children_col else None
            adult_quantity = parse_stage_quantity(row[adult_col]) if adult_col is not None and len(row) > adult_col else None

            stages.append(DeliveryStage(
                stage=stage,
                delivery_date=delivery_date,
                quantity=quantity,
                children_quantity=children_quantity,
                adult_quantity=adult_quantity,
            ))

    return stages


def summarize_delivery_stages(doc):
    years = set()
    count = 0

    for table in doc.tables:
        rows = [
            [normalize_text(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        header_index = find_delivery_header_index(rows)
        if header_index is None:
            continue

        header_rows = rows[header_index:header_index + 3]
        stage_col, date_col, _, _, _ = detect_delivery_columns(header_rows)
        if stage_col is None or date_col is None:
            continue

        for row in rows[header_index + 1:]:
            if len(row) <= max(stage_col, date_col):
                continue

            if not parse_stage(row[stage_col]):
                continue

            count += 1
            years.update(date.year for date in parse_all_dates(row[date_col]))

    return {
        "count": count,
        "years": sorted(years),
    }


def find_delivery_header_index(rows):
    for index, row in enumerate(rows):
        row_text = " ".join(row).lower()
        if is_delivery_header_text(row_text):
            return index
    return None


def normalize_key(value):
    return re.sub(r"[^а-яa-z0-9]", "", normalize_text(value).lower().replace("ё", "е"))


def is_delivery_header_text(row_text):
    row_key = normalize_key(row_text)
    has_stage = "этап" in row_key
    has_delivery = "поставк" in row_key or "отгрузк" in row_key
    has_deadline = "срок" in row_key or "дат" in row_key or "период" in row_key
    return has_stage and has_delivery and has_deadline


def detect_delivery_columns(header_rows):
    width = max(len(row) for row in header_rows)
    stage_col = date_col = children_col = adult_col = None
    quantity_cols = []

    for index in range(width):
        header = " ".join(
            row[index].lower()
            for row in header_rows
            if len(row) > index
        )
        header_key = normalize_key(header)
        if stage_col is None and "этап" in header_key:
            stage_col = index
        elif (
            date_col is None
            and ("срок" in header_key or "дат" in header_key or "период" in header_key)
            and ("поставк" in header_key or "отгрузк" in header_key)
        ):
            date_col = index
        elif children_col is None and "дети" in header_key:
            children_col = index
        elif adult_col is None and "взросл" in header_key:
            adult_col = index
        elif "количеств" in header_key or "объем" in header_key or "поставляем" in header_key:
            quantity_cols.append(index)

    quantity_cols = [
        index
        for index in quantity_cols
        if index not in {stage_col, date_col, children_col, adult_col}
    ]

    return stage_col, date_col, quantity_cols, children_col, adult_col


def parse_generic_stage_quantity(row, quantity_cols):
    quantities = []
    for column in quantity_cols:
        if len(row) <= column:
            continue
        quantity = parse_stage_quantity(row[column])
        if quantity is not None:
            quantities.append(quantity)

    if not quantities:
        return None

    total = sum(quantities)
    return int(total) if float(total).is_integer() else total


def parse_stage_quantity(value):
    text = normalize_text(value)
    if not text or text in {"-", "–"}:
        return None

    parenthesized = re.findall(r"\(([^()]*)\)", text)
    for part in reversed(parenthesized):
        quantity = parse_number(part)
        if quantity is not None:
            return quantity

    return parse_number(strip_quantity_footnote(text))


def strip_quantity_footnote(value):
    text = normalize_text(value)

    # In several Word templates footnote markers are glued to the number:
    # "977 2451" means "977 245" + footnote 1, not 9 772 451.
    if re.fullmatch(r"\d{1,3}(?:\s\d{3})+\d", text):
        return text[:-1]

    decimal_match = re.fullmatch(r"(\d[\d\s]*[,\.]\d{2})\d", text)
    if decimal_match:
        return decimal_match.group(1)

    return text


def parse_contract_unit_price(doc):
    for table in doc.tables:
        rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue

        table_text = " ".join(" ".join(row) for row in rows).lower()
        if "цена за единицу измерения товара" not in table_text:
            continue

        price_columns = detect_contract_unit_price_columns(rows[:3])
        if not price_columns:
            continue

        for row in rows[3:]:
            for column in price_columns:
                if len(row) <= column:
                    continue
                price = parse_number(row[column])
                if price is not None:
                    return price

    return None


def detect_contract_unit_price_columns(header_rows):
    width = max(len(row) for row in header_rows)
    result = []

    for index in range(width):
        header = " ".join(
            row[index].lower()
            for row in header_rows
            if len(row) > index
        )
        if "цена за единицу" in header and "итого" in header:
            result.append(index)

    if result:
        return result

    for index in range(width):
        header = " ".join(
            row[index].lower()
            for row in header_rows
            if len(row) > index
        )
        if "итого, руб" in header:
            result.append(index)

    return result


def parse_package_quantity(doc):
    for table in doc.tables:
        rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue

        table_text = " ".join(" ".join(row) for row in rows).lower()
        if "количество вторичных" not in table_text or "по потребности" not in table_text:
            continue

        package_columns = detect_package_quantity_columns(rows[:3])
        if not package_columns:
            continue

        total_row_value = None
        data_values = []
        for row in rows[3:]:
            row_text = " ".join(row).lower()
            row_values = []
            for column in package_columns:
                if len(row) <= column:
                    continue
                quantity = parse_stage_quantity(row[column])
                if quantity is not None:
                    row_values.append(quantity)

            if not row_values:
                continue

            row_sum = sum(row_values)
            if "итого" in row_text:
                total_row_value = row_sum
            else:
                data_values.append(row_sum)

        result = total_row_value if total_row_value is not None else sum(data_values) if data_values else None
        if result is not None:
            return int(result) if float(result).is_integer() else result

    return None


def detect_package_quantity_columns(header_rows):
    width = max(len(row) for row in header_rows)
    result = []

    for index in range(width):
        header = " ".join(
            row[index].lower()
            for row in header_rows
            if len(row) > index
        )
        if "количество вторичных" in header and "по потребности" in header:
            result.append(index)

    return result


def parse_document_deadlines(doc):
    clause = find_document_deadline_clause(doc)
    if not clause:
        return {}

    deadlines = {}
    stage_deadlines = parse_stage_deadlines(clause)
    if stage_deadlines:
        deadlines.update(stage_deadlines)
        return deadlines

    dates = parse_all_dates(clause)
    if dates:
        deadlines["I"] = dates[-1]

    return deadlines


def find_document_deadline_clause(doc):
    for clause_number in ("9.4.", "9.3.", "9.5.", "10.1.", "10.2."):
        clause = extract_clause_text(doc, clause_number)
        if is_document_deadline_text(clause):
            return clause

    candidates = []
    for line in collect_document_lines(doc):
        line_lower = normalize_text(line).lower()
        if is_document_deadline_text(line_lower) and parse_all_dates(line_lower):
            candidates.append(line)

    return " ".join(candidates)


def is_document_deadline_text(text):
    text = normalize_text(text).lower()
    if not text:
        return False

    has_document = "документ" in text or "еис" in text or "приемк" in text
    has_action = "предостав" in text or "размещ" in text or "формир" in text or "направ" in text
    return has_document and (has_action or bool(parse_all_dates(text)))


def parse_stage_deadlines(text):
    deadlines = parse_date_before_stage_deadlines(text)
    stage_matches = list(re.finditer(
        r"(?:по|для|за)?\s*"
        r"(перв\w*|втор\w*|трет\w*|четвер\w*|пят\w*|шест\w*|i{1,3}|iv|v|vi|1|2|3|4|5|6)\s+этап",
        text,
        flags=re.IGNORECASE,
    ))

    for index, match in enumerate(stage_matches):
        stage = parse_stage(match.group(0))
        if not stage:
            continue
        if stage in deadlines:
            continue

        next_start = stage_matches[index + 1].start() if index + 1 < len(stage_matches) else len(text)
        after_segment = text[match.end():next_start]
        dates = parse_all_dates(after_segment)
        if dates:
            deadlines[stage] = dates[-1]

    return deadlines


def parse_date_before_stage_deadlines(text):
    deadlines = {}
    text = normalize_text(text)

    for parsed_date, _, date_end in parse_all_date_matches(text):
        tail = text[date_end:date_end + 180]
        match = re.search(
            r"^\s*(?:г\.)?.{0,160}?[-–—]\s*"
            r"(?:по|для|за)?\s*"
            r"(перв\w*|втор\w*|трет\w*|четвер\w*|пят\w*|шест\w*|i{1,3}|iv|v|vi|1|2|3|4|5|6)\s+этап",
            tail,
            flags=re.IGNORECASE,
        )
        if not match:
            continue

        stage = parse_stage(match.group(0))
        if stage:
            deadlines[stage] = parsed_date

    return deadlines


def parse_all_date_matches(value):
    text = normalize_text(value).lower()
    matches = []

    for match in re.finditer(r"(\d{2})\.(\d{2})\.(\d{4})", text):
        day, month, year = map(int, match.groups())
        matches.append((date(year, month, day), match.start(), match.end()))

    russian_date_pattern = (
        r"[«\"]?(\d{1,2})[»\"]?\s+"
        r"(января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря)"
        r"\s+(\d{4})"
    )
    for match in re.finditer(russian_date_pattern, text):
        day = int(match.group(1))
        month = MONTHS[match.group(2)]
        year = int(match.group(3))
        matches.append((date(year, month, day), match.start(), match.end()))

    return sorted(matches, key=lambda item: item[1])


def extract_clause_text(doc, clause_number):
    lines = collect_document_lines(doc)
    clause_lines = []
    collecting = False
    clause_prefix = clause_number.rstrip(".")

    for line in lines:
        stripped = line.lstrip("«\"")

        if not collecting:
            if stripped.startswith(clause_number):
                clause_lines.append(stripped)
                collecting = True
            continue

        clause_match = re.match(r"(\d+(?:\.\d+)+)\.", stripped)
        if clause_match:
            current_clause = clause_match.group(1)
            if current_clause != clause_prefix and not current_clause.startswith(f"{clause_prefix}."):
                break

        clause_lines.append(stripped)

    return " ".join(clause_lines)
